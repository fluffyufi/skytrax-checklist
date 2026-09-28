"""Quick self-test: fill scenario inputs into a copy, recalc with LibreOffice, print key results."""
import shutil
import subprocess
import sys
from datetime import datetime

import openpyxl

SRC = "/home/user/skytrax-checklist/output/MAGCS_Skytrax_2026_Catering_Readiness.xlsx"
DST = "/tmp/claude-0/-home-user-skytrax-checklist/c0ab9c35-2f2c-5eff-909f-a4943aca1b4a/scratchpad/selftest.xlsx"
RECALC = "/root/.claude/skills/synced/dbf11ce3-ba00-4160-a445-f73a3231a942_7e2a8d9e-80e6-4124-9d8d-d105bdc3f685/xlsx/scripts/recalc.py"

shutil.copy(SRC, DST)
wb = openpyxl.load_workbook(DST)
ws = wb["Checks"]
wb["Settings"]["B4"] = datetime(2026, 10, 9, 12, 0)  # as-of UTC
rows = {ws.cell(r, 1).value: r for r in range(5, ws.max_row + 1) if ws.cell(r, 1).value}


def put(cid, **kw):
    r = rows[cid]
    for col, v in kw.items():
        ws[f"{col}{r}"] = v


# F02 T-24 pass without batch -> INVALID; F02 T7-01 full pass -> COMPLETE
put("F02-T7-01", U="Pass", V="ok", Z="rep1", AD=datetime(2026, 10, 2, 10, 0), AE="Ver")
put("F02-T24-01", U="Pass", V="ok", Z="rep", AD=datetime(2026, 10, 8, 22, 0), AE="Ver")
put("F02-T7-02", U="Pass")  # INVALID missing
put("F02-T7-04", U="N/A")  # INVALID no justification
put("F02-T7-05", U="N/A", AC="just", AE="Ver")  # justified
put("F02-T12-02", U="Pass", X=100, Y=98, Z="e", AD=datetime(2026, 10, 9, 9, 0), AE="v")  # variance invalid
wb.save(DST)
print(subprocess.run([sys.executable, RECALC, DST, "300"], capture_output=True, text=True).stdout[-600:])
wb = openpyxl.load_workbook(DST, data_only=True)
ws = wb["Checks"]
for cid in ("F02-T7-01", "F02-T7-02", "F02-T7-03", "F02-T7-04", "F02-T7-05", "F02-T24-01", "F02-T12-02",
            "F12-T7-03", "F21-T24-01", "F21-T12-01", "F07-T12-09", "F01-T7-01"):
    r = rows[cid]
    print(cid, ws[f"G{r}"].value, "| N:", ws[f"N{r}"].value, "| due UTC", ws[f"Q{r}"].value, "| local", ws[f"R{r}"].value,
          "| state:", ws[f"AH{r}"].value, ws[f"AI{r}"].value, ws[f"AJ{r}"].value, ws[f"AK{r}"].value, ws[f"AL{r}"].value)
fl = wb["Flights"]
for r in (5, 6, 7, 15, 17, 19, 25):
    print([fl.cell(r, c).value for c in (1, 6, 9, 26, 27, 28, 29, 31, 33, 35, 36, 40, 41, 42, 43, 44, 45, 46)])

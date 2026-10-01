"""Fill flight F02 (MH0727 KUL-CGK) completely & validly, confirm READY, then break items one at a time."""
import shutil, subprocess, sys
from datetime import datetime, timedelta
import openpyxl

import os
SRC = os.environ.get("TEST_SRC", "/home/user/skytrax-checklist/output/MAGCS_Skytrax_2026_Catering_Readiness.xlsx")
TMP = "/tmp/claude-0/-home-user-skytrax-checklist/c0ab9c35-2f2c-5eff-909f-a4943aca1b4a/scratchpad/ready_{}.xlsx"
RECALC = "/root/.claude/skills/synced/dbf11ce3-ba00-4160-a445-f73a3231a942_7e2a8d9e-80e6-4124-9d8d-d105bdc3f685/xlsx/scripts/recalc.py"

# a killed LibreOffice run leaves '.~lock.<file>#' behind; the next recalc of that file then stalls on a read-only prompt
import glob as _glob
for _lk in _glob.glob(os.path.join(os.path.dirname(TMP), ".~lock.*#")):
    try:
        os.remove(_lk)
    except OSError:
        pass
STD = datetime(2026, 10, 9, 21, 50)  # local KUL
TAIL = {"B738MAX": "9M-MVO", "A333": "9M-MTJ", "A350": "9M-MAB", "A339": "9M-MTA"}


def fill(ws, docs, fid="F02"):
    for r in range(5, ws.max_row + 1):
        if ws[f"B{r}"].value != fid:
            continue
        cid, cp, item, cat = ws[f"A{r}"].value, ws[f"G{r}"].value, ws[f"J{r}"].value, ws[f"I{r}"].value
        if isinstance(ws[f"N{r}"].value, str) and ws[f"N{r}"].value.startswith("=IF(TRIM(Flights!$L"):
            continue  # IFE rule row
        t = {"T-7D": STD - timedelta(days=7, hours=2), "T-24H": STD - timedelta(hours=26),
             "T-12H PREP": STD - timedelta(hours=13), "UPLIFT": STD - timedelta(hours=2)}[cp]
        ws[f"T{r}"] = "A. Rahman"
        ws[f"V{r}"] = "Checked against menu checklist, all to spec"
        ws[f"S{r}"] = "Form / checklist"
        ws[f"Z{r}"] = f"QF-{1000 + r}"
        ws[f"AE{r}"] = "N. Ismail"
        ws[f"AD{r}"] = t
        if ws[f"AQ{r}"].value == 1:
            ws[f"W{r}"] = "PASB-261008-BC-017"
        if ws[f"AR{r}"].value == 1:
            ws[f"X{r}"] = 12
            ws[f"Y{r}"] = 12
        if ws[f"N{r}"].value == "Clarification required":
            ws[f"U{r}"] = "N/A"
            ws[f"AC{r}"] = "Confirmed – not carried on this sector"
            ws[f"S{r}"] = "Email"; ws[f"Z{r}"] = f"EM-26-{1000 + r}"
        else:
            ws[f"U{r}"] = "Pass"
    d = docs
    fr = int(fid[1:]) + 4
    d[f"F{fr}"] = f"https://magcs.sharepoint.com/catering/GLD/{fid}-GLD.pdf"
    d[f"H{fr}"] = f"https://magcs.sharepoint.com/catering/MCL/{fid}-menu-checklist.pdf"
    fl = d.parent["Flights"]
    fl[f"N{fr}"] = TAIL.get(fl[f"L{fr}"].value, "9M-MTA")


def run(tag, mutate=None, asof=datetime(2026, 10, 9, 13, 49)):
    p = TMP.format(tag)
    shutil.copy(SRC, p)
    wb = openpyxl.load_workbook(p)
    wb["Settings"]["B4"] = asof
    fill(wb["Checks"], wb["Documents"])
    rows = {wb["Checks"].cell(r, 1).value: r for r in range(5, wb["Checks"].max_row + 1)}
    if mutate:
        mutate(wb, rows)
    wb.save(p)
    out = subprocess.run([sys.executable, RECALC, p, "900"], capture_output=True, text=True).stdout
    wb = openpyxl.load_workbook(p, data_only=True)
    fl = wb["Flights"]
    bad = [(wb["Checks"][f"A{r}"].value, wb["Checks"][f"AH{r}"].value) for r in range(5, wb["Checks"].max_row + 1)
           if wb["Checks"][f"B{r}"].value == "F02" and not str(wb["Checks"][f"AH{r}"].value).startswith(("COMPLETE", "N/A"))]
    ok = '"total_errors": 0' in out
    print(f"{tag:28s} readiness={fl['AT6'].value!r:40s} overall={fl['AN6'].value} recalc_ok={ok} bad={bad[:3]}")


def m(col_vals):
    def f(wb, rows):
        for cid, col, v in col_vals:
            wb["Checks"][f"{col}{rows[cid]}"] = v
    return f


if __name__ == "__main__":
  run("baseline_full")
  run("after_departure_T24", m([("F02-T24-01", "AD", datetime(2026, 10, 9, 21, 55))]), asof=datetime(2026, 10, 9, 14, 0))
  run("pic_missing", m([("F02-T7-01", "T", None)]))
  run("whitespace_evidence", m([("F02-T7-02", "Z", " ")]))
  run("na_on_meal_uplift", m([("F02-UPL-01", "U", "N/A"), ("F02-UPL-01", "AC", "not needed")]))
  run("status_trailing_space", m([("F02-T7-01", "U", "Pass ")]))
  run("uplift_before_window", m([("F02-UPL-02", "AD", datetime(2026, 10, 9, 14, 0))]))
  run("qty_variance", m([("F02-T12-02", "Y", 11)]))
  run("uplift_blank_not_due", m([("F02-UPL-01", "U", None)]))
  run("doc_outstanding", lambda wb, rows: wb["Documents"].__setitem__("H6", None))
  run("doc_placeholder", lambda wb, rows: wb["Documents"].__setitem__("F6", "TBC"))
  run("doc_named_hyperlink_text", lambda wb, rows: wb["Documents"].__setitem__("H6", "Menu Checklist MH0727 YCL"))
  run("na_on_required_line", m([("F02-T12-12", "U", "N/A"), ("F02-T12-12", "AC", "Confirmed – not carried on this sector")]))
  run("tail_missing", lambda wb, rows: wb["Flights"].__setitem__("N6", None))
  run("tail_not_in_group", lambda wb, rows: wb["Flights"].__setitem__("N6", "9M-MXE"))
  run("tail_in_group", lambda wb, rows: wb["Flights"].__setitem__("N6", "9M-MVR"))
  run("r21_one_bad_uplift_time", m([("F02-UPL-02", "AD", STD - timedelta(hours=20))]))

  def uplift_blank(wb, rows):
      ck = wb["Checks"]
      for c, r in rows.items():
          if c and c.startswith("F02-UPL-"):
              for k in ("U", "AD"):
                  ck[f"{k}{r}"] = None
  run("label_prep_done_awaiting_uplift", uplift_blank, asof=datetime(2026, 10, 9, 10, 50))

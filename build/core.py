"""Core sheets: Settings, Flights, Checks, Documents, Requirements, Instructions.

Every due time, completion %, state and readiness value is a live formula.
See CONTRACT.md for the column map shared with dashboard.py / printable.py.
"""
from datetime import datetime, timedelta

from openpyxl.comments import Comment
from openpyxl.formatting.rule import CellIsRule, FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.datavalidation import DataValidation

NAVY = "1F3864"
F_HEAD = PatternFill("solid", fgColor=NAVY)
F_INPUT = PatternFill("solid", fgColor="FFF2CC")
F_CALC = PatternFill("solid", fgColor="F2F2F2")
F_REF = PatternFill("solid", fgColor="E8EEF7")
FONT = "Arial"
THIN = Side(style="thin", color="BFBFBF")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
WRAP = Alignment(wrap_text=True, vertical="top")
CENTER = Alignment(horizontal="center", vertical="top", wrap_text=True)
DT = "dd-mmm-yy hh:mm"
ND = "–"  # en dash used in all state strings

STATUS_LIST = ["Not started", "In progress", "Pass", "Pass after CA", "Fail", "N/A"]
CA_LIST = ["Open", "Closed"]


def f(size=10, bold=False, color="000000", italic=False):
    return Font(name=FONT, size=size, bold=bold, color=color, italic=italic)


def header(ws, row, labels, widths=None, height=42):
    for i, lab in enumerate(labels, 1):
        c = ws.cell(row, i, lab)
        c.font = f(9, True, "FFFFFF")
        c.fill = F_HEAD
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        c.border = BORDER
    ws.row_dimensions[row].height = height
    if widths:
        for i, w in enumerate(widths, 1):
            ws.column_dimensions[get_column_letter(i)].width = w


def title(ws, text, sub=None, sub_span=None, sub_width=None):
    """Short title in A1 (must fit the repeated print-title columns); subtitle wrapped inside sub_span on row 2."""
    ws["A1"] = text
    ws["A1"].font = f(14, True, NAVY)
    ws.row_dimensions[1].height = 24
    if sub:
        anchor = sub_span.split(":")[0] if sub_span else "A2"
        ws[anchor] = sub
        ws[anchor].font = f(9, False, "595959", True)
        if sub_span:
            ws.merge_cells(sub_span)
            ws[anchor].alignment = Alignment(wrap_text=True, vertical="top")
            ws.row_dimensions[2].height = 12.5 * est_lines(sub, sub_width) + 4


def est_lines(text, width):
    if text is None:
        return 1
    chars_per_line = max(1, int(width * 1.15))
    n = 0
    for part in str(text).split("\n"):
        n += max(1, -(-len(part) // chars_per_line))
    return n


def off(stn, utc):
    """UTC offset (hours) of station `stn` at UTC instant `utc` — both formula fragments."""
    m = f"MATCH({stn},TZ_STN,0)"
    return (f"IF(AND(INDEX(TZ_DSTS,{m})<>\"\",{utc}>=INDEX(TZ_DSTS,{m}),{utc}<INDEX(TZ_DSTE,{m})),"
            f"INDEX(TZ_DST,{m}),INDEX(TZ_STD,{m}))")


def stdoff(stn):
    return f"INDEX(TZ_STD,MATCH({stn},TZ_STN,0))"


def dt(s):
    return datetime.strptime(s, "%Y-%m-%d %H:%M")


def name(wb, nm, ref):
    wb.defined_names[nm] = DefinedName(nm, attr_text=ref)


# ------------------------------------------------------------------ Settings
SQ_FIRST_ROW = 27  # Settings: station-confirmation table data rows start here (header on 26)


def station_questions(data):
    """Open 'where is it loaded' questions: prep line + on-board line whose uplift station the reference leaves open.
    Returns one dict per question with its Settings row, the two possible stations and the shortest agenda block
    from the alternative station to the departure station (for the carrying flight's arrival)."""
    fl = {f["id"]: f for f in data["flights"]}
    out = []
    for ck in data["checks"]:
        f = fl[ck["flight_id"]]
        if ck["check_type"] != "Preparation" or ck["applic"] != "Clarification required":
            continue
        if ck["item"].startswith("Toiletry kits") and ck["uplift_stn"] == "LHR or KUL - confirm":
            alt = "LHR"
        elif ck["item"].startswith("Sales cart") and f["dep"] != "KUL" and ck["uplift_stn"] == "Not stated in reference":
            alt = "KUL"
        else:
            continue
        upl = next(x["check_id"] for x in data["checks"] if x.get("prep_link") == ck["check_id"])
        blocks = [g["block_h"] for g in data["flights"] if g["dep"] == alt and g["arr"] == f["dep"]]
        out.append(dict(qid=f"SQ{len(out) + 1}", flight=f["id"], flt=f["flt"], item=ck["item"].replace(" - prepared", ""),
                        prep=ck["check_id"], upl=upl, dep=f["dep"], default=ck["station"], alt=alt,
                        block=min(blocks), row=SQ_FIRST_ROW + len(out)))
    return out


def build_settings(wb, data):
    ws = wb.create_sheet("Settings")
    title(ws, "Settings & reference tables", "Yellow cells are editable. Everything else is referenced by formulas.")
    rows = [
        (4, "As-of override (UTC). Leave blank for the live clock; enter a UTC date-time to review the workbook as at that moment.", None, True),
        (5, "This PC's clock offset from UTC (hours). 8 = Malaysia (MYT). Used to convert NOW() to UTC.", 8, True),
        (6, "Effective as-of time (UTC) used for all overdue tests", "=IF(ISNUMBER(B4),B4,NOW()-N(B5)/24)", False),
        (7, "Physical-uplift window: on-board confirmation is valid only within this many hours before STD, and not before a carrying flight could have arrived (assumption – adjust to station loading practice; max 10)", 6, True),
        (8, "T-24H sensory test: earliest valid completion, hours before its due time (assumption – production batch must exist)", 24, True),
        (9, "T-12H preparation check: earliest valid completion, hours before its due time (assumption)", 12, True),
        (10, "T-7D checks: earliest valid completion, days before the T-7D due time (assumption – older evidence is stale)", 14, True),
    ]
    ws.column_dimensions["A"].width = 50
    ws.column_dimensions["B"].width = 20
    for r, lab, val, inp in rows:
        ws.cell(r, 1, lab).font = f(9)
        ws.cell(r, 1).alignment = WRAP
        c = ws.cell(r, 2, val)
        c.font = f(10, True, "0000FF" if inp else "000000")
        c.fill = F_INPUT if inp else F_CALC
        c.border = BORDER
        ws.row_dimensions[r].height = max(26, 12 * est_lines(lab, 50) + 4)
    ws["B4"].number_format = DT
    ws["B6"].number_format = DT
    ws["C6"] = "=B6+8/24"
    ws["C6"].number_format = DT
    ws["D6"] = "← as-of in MYT"
    ws["D6"].font = f(8, italic=True, color="595959")
    ws["B4"].comment = Comment("Blank = live. Example: 2026-10-09 12:00 to see what is overdue at that UTC time.", "MAGCS")
    name(wb, "AsOfUTC", "Settings!$B$6")
    for rr, lo, hi, dflt in ((7, 1, 10, 6), (8, 1, 72, 24), (9, 1, 48, 12), (10, 1, 60, 14)):
        ws[f"C{rr}"] = f"=IF(ISNUMBER(B{rr}),MIN(MAX(B{rr},{lo}),{hi}),{dflt})"
        ws[f"C{rr}"].font = f(10, True)
        ws[f"C{rr}"].fill = F_CALC
        ws[f"C{rr}"].border = BORDER
        ws[f"D{rr}"] = f"← value in use (limited to {lo}–{hi})"
        ws[f"D{rr}"].font = f(8, italic=True, color="595959")
    name(wb, "UpliftWindowH", "Settings!$C$7")
    name(wb, "T24EarlyH", "Settings!$C$8")
    name(wb, "T12EarlyH", "Settings!$C$9")
    name(wb, "T7EarlyD", "Settings!$C$10")
    ws["K3"] = "Placeholder words rejected (compared after removing spaces & punctuation, any case)"
    ws["K3"].font = f(9, True, NAVY)
    ph = ["na", "tbc", "tba", "tbd", "tbconfirmed", "tobeconfirmed", "tobeadvised", "pending", "awaiting", "none",
          "nil", "null", "same", "later", "unknown", "x", "xx", "xxx", "test", "dummy", "notapplicable", "nophoto", "noevidence",
          "visualcheck", "visuallychecked", "visuallyinspected", "checkedok", "sighted", "confirmed", "notrequired",
          "phototaken", "byphone", "aspergld", "attached", "visually", "self", "sameaspic", "pic", "me", "myself",
          "good", "passed", "satisfactory", "accepted", "fine", "noted", "complete", "completed", "pass", "allgood",
          "allok", "allfine", "aspermenu", "asperspec", "asrequired", "asperstandard", "conforms", "conform",
          "acceptable", "verified", "checkedandverified", "allcorrect", "correct", "okay", "noissues", "noissue",
          "allinorder", "inorder", "tospec", "allsatisfactory", "goodcondition", "ingoodcondition",
          "notavailable", "notprovided", "nobatch", "nobatchno", "nobatchnumber", "seelabel", "unknownbatch",
          "nonerequired", "nocarequired", "noca", "nilca", "noactionrequired", "noactionneeded", "nonenecessary",
          "qa", "qaqc", "csm", "dutycsm", "supervisor", "caterer", "pasb", "mcat", "crew", "team", "cic", "ccic",
          "manager", "officer", "qaofficer", "station", "stationstaff", "staff", "inspector",
          "done", "yes", "checked", "ok", "okay", "visual", "verbal", "naverbal", "photo", "email", "seen", "fine"]
    for i, v in enumerate(ph):
        ws.cell(4 + i, 11, v).font = f(9)
    ws.column_dimensions["K"].width = 22
    name(wb, "L_Placeholder", f"Settings!$K$4:$K${3 + len(ph)}")
    # 'still pending' phrases (whole words) – rejected in result, batch, N/A justification, evidence and attachments
    pend = ["tbc", "tbd", "tba", "pending", "awaiting", "waiting", "to follow", "will follow", "to be confirmed",
            "to be advised", "not yet", "yet to", "not confirmed", "not received", "to be provided", "will be uploaded",
            "will be emailed", "will be sent", "will be attached", "will be provided", "to be attached", "to be sent",
            "to be emailed", "to be uploaded", "not uploaded", "will upload", "will email", "will send", "upload later",
            "send later"]
    # ...each occurrence cancelled by a negated form ("nothing pending", "no CA pending", "not awaiting anything")
    objs = ["ca", "cas", "corrective action", "action", "actions", "item", "items", "issue", "issues", "clarification",
            "clarifications", "document", "documents", "answer", "reply", "query", "queries", "further action", "follow up"]
    single = [w for w in pend if " " not in w]
    pneg = sorted({f"{p} {w}" for w in single for p in ("no", "nothing", "not", "none", "nil")} |
                  {f"no {o} {w}" for o in objs for w in single} |
                  {"nothing to follow", "nothing further to follow", "nothing more to follow", "nothing to be confirmed",
                   "nothing to be sent", "nothing to be uploaded", "not awaiting", "no longer awaiting", "no longer pending",
                   "not yet effective", "not yet introduced", "not yet applicable", "not yet in service",
                   "yet to be launched", "yet to be introduced", "or to follow", "or pending", "nothing awaited",
                   "no pending", "no pending items", "nothing outstanding or to follow", "awaiting pax",
                   "awaiting pax boarding", "awaiting boarding", "awaiting passengers", "pending nil", "pending none",
                   "pending items nil", "pending items none", "pending 0", "pending items 0", "awaiting nothing"})
    # preparation-stage wording: cancels 'pending' ONLY on Preparation rows (T-12H PREP is before loading by design)
    pprep = sorted({"awaiting loading", "awaiting uplift", "awaiting dispatch", "awaiting transport", "pending loading",
                   "pending uplift", "pending dispatch", "not yet loaded", "not yet uplifted", "not yet dispatched",
                   "yet to be loaded", "yet to be uplifted", "yet to be dispatched", "waiting in", "waiting for loading",
                   "waiting for uplift", "waiting for dispatch",
                   "awaiting delivery", "awaiting collection", "awaiting truck", "awaiting pickup", "awaiting pick up",
                   "awaiting aircraft", "awaiting boarding", "pending transport", "pending delivery", "pending aircraft",
                   "pending collection", "waiting at", "waiting for truck", "waiting for collection", "not yet boarded",
                   "not yet arrived"})
    # on-board rows: wording that says the item is NOT on board blocks a physical-uplift confirmation
    notload = ["not loaded", "not on board", "not onboard", "not uplifted", "not yet loaded", "not yet uplifted",
               "not yet on board", "yet to be loaded", "yet to be uplifted", "awaiting loading", "awaiting uplift",
               "pending loading", "pending uplift", "not delivered", "left behind", "offloaded", "off loaded",
               "to be loaded", "will be loaded", "to be uplifted", "will be uplifted", "ready for loading",
               "ready for uplift", "loading in progress", "being loaded", "still loading", "loading after",
               "delivered to aircraft side", "at aircraft side"]
    # ...cancelled when negated ("nothing left behind", "0 offloaded", "none offloaded")
    nl_neg = sorted({f"{p} {w}" for w in ("left behind", "offloaded", "off loaded", "not loaded", "not on board",
                                             "not uplifted", "not delivered")
                     for p in ("no", "nothing", "none", "0", "zero", "nil", "no items", "no meals")} |
                    {f"{w} {q}" for w in ("offloaded", "left behind", "not loaded", "not on board")
                     for q in ("0", "nil", "none")} |
                    {"aircraft side and loaded", "aircraft side and uplifted", "aircraft side then loaded",
                     "aircraft side then stowed", "aircraft side and stowed",
                     "at aircraft side and loaded"})
    # strict list for evidence, N/A justifications and document attachments: any future / not-yet wording
    strict = pend + ["awaited", "to be forwarded", "will be forwarded", "to be scanned", "will be scanned", "to be filed",
                     "will be filed", "to be shared", "will be shared", "will share", "will forward", "not sent",
                     "not shared", "not forwarded", "not scanned", "not filed", "still with", "hard copy outstanding",
                     "evidence outstanding", "copy outstanding", "documents outstanding", "being prepared", "to come",
                     "once received", "requested from", "to upload", "draft only", "not in yet", "will be sent tomorrow",
                     "expected tomorrow", "due tomorrow"]
    # extra phrases rejected only as evidence / attachment (the record itself is not on file)
    ev = ["as above", "see above", "ditto", "verbal", "verbally", "refer above", "same as above", "by phone",
          "phone call", "on the phone", "over the phone", "told", "told by", "call with", "whatsapp call", "by call",
          "telephone", "telephoned", "phoned", "called", "via phone", "phone confirmation", "rang", "orally",
          "oral confirmation", "tel call", "tel", "on call", "on a call", "via call"]
    # whole words meaning a result is not a plain pass
    rb = ["reject", "rejected", "rejects", "fail", "failed", "fails", "failure", "not ok", "nok", "unsatisfactory",
          "not acceptable", "unacceptable", "not satisfactory", "below standard", "shortfall", "shortage", "short by",
          "discrepancy", "discrepancies", "defect", "defective", "defects", "dirty", "damaged", "missing", "broken",
          "leaking", "expired", "non conforming", "nonconforming", "not to spec", "not delivered", "not received",
          "not available", "not provided", "not supplied", "not loaded", "not on board", "incorrect", "wrong",
          "stained", "torn", "spoiled", "spoilt", "contaminated", "insufficient", "not enough", "unavailable",
          "out of stock", "short", "late", "undelivered", "mismatch", "not matching", "does not match",
          "not match", "wrongly", "mislabelled", "mislabeled", "unlabelled", "unlabeled", "partially", "partial",
          "incomplete", "not complete", "not completed", "except", "but no", "but not", "out of date", "mouldy",
          "moldy", "mould", "mold", "inoperative", "not working", "unserviceable", "spilled", "spilt", "spillage",
          "leak", "leaked", "not all", "not prepared", "too high", "too low", "too warm", "too cold", "left at",
          "missed", "not boarded", "after departure", "arrived late", "not ready", "unfit", "not assigned",
          "not acknowledged", "unacknowledged", "outstanding", "still at", "cracked", "unusable", "too salty",
          "too dry", "too sweet", "too spicy", "overcooked", "undercooked", "re cook", "recook", "redo",
          "re do", "fewer", "fewer than", "not sealed", "unsealed", "below required", "below count", "sent back",
          "rework", "cold on arrival", "missing from", "superseded", "previous cycle", "old cycle", "doors closed",
          "exceeds", "exceeded", "out of spec", "outside spec", "outside limits", "above limit", "over limit",
          "instead of", "in place of", "substituted with", "replaced with"] + \
         [f"only {n}" for n in range(1, 100)] + \
         [f"short {n}" for n in range(1, 31)] + [f"{n} short" for n in range(1, 31)]
    # ...but not when that same word is negated or zero-counted (each phrase below cancels one occurrence)
    pre = ["no", "nil", "zero", "0", "nothing", "not", "without", "none", "free of", "no sign of", "no evidence of",
           "or", "nor",  # 'or'/'nor' carry a negation across a list: "no stains or defects"
           "no visible", "no quality", "no quantity", "no qty", "no meal", "no loading", "no printing", "no equipment",
           "no cleanliness", "no presentation", "no obvious", "any", "free from", "no items", "no meals", "no trays"]
    post = ["0", "nil", "none", "zero", "free", "nothing", "not found", "found none", "found 0", "found nil", "noted 0"]
    base = [w for w in rb if not w[0].isdigit() and not w.startswith(("short ", "only "))]
    negp = sorted({f"{p} {w}" for w in base for p in pre} | {f"{w} {q}" for w in base for q in post} |
                  {"fail safe", "failsafe", "short rib", "short ribs", "short grain", "short crust", "short haul",
                   "short sector", "short pasta", "short cake", "short term", "short notice", "broken rice",
                   "dirty linen", "dirty trays", "dirty tray", "for dirty", "late night", "late supper",
                   "except none", "mold free", "mould free", "no exceptions", "leak test", "leak check",
                   "leak tested", "damaged cart log", "damage log", "missed nothing", "nothing missed",
                   "not all required", "none left at"} |
                  {f"{w} items {q}" for w in base for q in ("none", "nil", "0", "zero")})
    for col, title_txt, items in (("L", "Evidence phrases that point elsewhere instead of to a record", ev),
                                  ("M", "Result words that cannot be a plain Pass (whole words)", rb),
                                  ("N", "…cancelled when negated / zero-counted (per word)", negp),
                                  ("O", "'Still pending' phrases – result, evidence, batch, N/A justification, attachments", pend),
                                  ("P", "…cancelled when negated (per occurrence)", pneg),
                                  ("R", "…also cancelled on Preparation rows only (prep-stage wording)", pprep),
                                  ("S", "Blocks an on-board (physical uplift) confirmation", notload),
                                  ("T", "…cancelled when negated", nl_neg),
                                  ("Q", "Strict 'not on file yet' words – evidence, N/A justification, attachments", strict)):
        ws[f"{col}3"] = title_txt
        ws[f"{col}3"].font = f(9, True, NAVY)
        ws[f"{col}3"].alignment = WRAP
        ws.column_dimensions[col].width = 22
        for i, v in enumerate(items):
            ws.cell(4 + i, ord(col) - 64, v).font = f(9)
    ws.row_dimensions[3].height = 48
    name(wb, "L_EvidencePhrase", f"Settings!$L$4:$L${3 + len(ev)}")
    name(wb, "L_ResultBad", f"Settings!$M$4:$M${3 + len(rb)}")
    name(wb, "L_NegPhrase", f"Settings!$N$4:$N${3 + len(negp)}")
    name(wb, "L_Pending", f"Settings!$O$4:$O${3 + len(pend)}")
    name(wb, "L_PendNeg", f"Settings!$P$4:$P${3 + len(pneg)}")
    name(wb, "L_PendPrep", f"Settings!$R$4:$R${3 + len(pprep)}")
    name(wb, "L_NotLoaded", f"Settings!$S$4:$S${3 + len(notload)}")
    name(wb, "L_NotLoadedNeg", f"Settings!$T$4:$T${3 + len(nl_neg)}")
    ws["U3"] = "Evidence types (Checks column S)"
    ws["U3"].font = f(9, True, NAVY)
    ws["U3"].alignment = WRAP
    ws.column_dimensions["U"].width = 26
    for i, v in enumerate(EVIDENCE_TYPES):
        ws.cell(4 + i, 21, v).font = f(9)
    name(wb, "L_EvidenceType", f"Settings!$U$4:$U${3 + len(EVIDENCE_TYPES)}")
    # ---------------- station confirmations: where an item with an open uplift station is actually loaded
    sq = station_questions(data)
    hdr = SQ_FIRST_ROW - 1
    ws.cell(hdr - 1, 1, "STATION CONFIRMATIONS – where is the item loaded? (answer from MAGCS / caterer; yellow = input)").font = f(10, True, NAVY)
    for c, h in enumerate(["Question / flight", "Item", "Lines", "Departure stn", "Other possible stn",
                           "CONFIRMED uplift stn", "Carrying flight departs that stn (UTC)", "Shortest block (h)",
                           "Carrier arrives (UTC)", "Status"], 1):
        cell = ws.cell(hdr, c, h)
        cell.font = f(9, True, "FFFFFF"); cell.fill = PatternFill("solid", fgColor=NAVY); cell.alignment = WRAP; cell.border = BORDER
    ws.row_dimensions[hdr].height = 40
    for q in sq:
        r = q["row"]
        vals = [f"{q['qid']}  {q['flight']} {q['flt']}", q["item"], f"{q['prep']} / {q['upl']}", q["dep"], q["alt"], None, None,
                round(q["block"], 2), f"=IF(ISNUMBER(G{r}),G{r}+H{r}/24,\"\")",
                (f"=IF(TRIM(F{r})=\"\",\"OPEN – confirm the uplift stn (reference default: {q['default']})\","
                 f"IF(TRIM(F{r})=D{r},\"Loaded at departure stn\",IF(ISNUMBER(G{r}),\"Loaded at \"&F{r}&\" – carried on the inbound flight\","
                 f"\"Enter the carrying flight's departure from \"&F{r}&\" (UTC)\")))")]
        for c, v in enumerate(vals, 1):
            cell = ws.cell(r, c, v)
            cell.font = f(9); cell.border = BORDER; cell.alignment = WRAP
        for c in (6, 7):
            ws.cell(r, c).fill = F_INPUT
        ws.cell(r, 7).number_format = DT
        ws.cell(r, 9).number_format = DT
        dvs = DataValidation(type="list", formula1=f'"{q["dep"]},{q["alt"]}"', allow_blank=True, showErrorMessage=True,
                             errorTitle="Uplift station", error=f"Choose {q['dep']} or {q['alt']}.")
        dvs.add(f"F{r}")
        ws.add_data_validation(dvs)
        dvg = DataValidation(type="decimal", operator="between", formula1="46204", formula2="46419", allow_blank=True,
                             showErrorMessage=True, error="Enter the carrying flight's departure as a UTC date-time.")
        dvg.add(f"G{r}")
        ws.add_data_validation(dvg)
        ws.row_dimensions[r].height = 30

    for col, title_txt, items, nm in (("X", "Outcome (Checks AC) – clarification lines and N/A", OUTCOMES, "L_Outcome"),
                                      ("V", "Role / title words – not a person's name (PIC, verifier)", ROLE_WORDS, "L_Role"),):
        ws[f"{col}3"] = title_txt
        ws[f"{col}3"].font = f(9, True, NAVY)
        ws[f"{col}3"].alignment = WRAP
        ws.column_dimensions[col].width = 22
        for i, v in enumerate(items):
            ws.cell(4 + i, ord(col) - 64, v).font = f(9)
        name(wb, nm, f"Settings!${col}$4:${col}${3 + len(items)}")
    name(wb, "L_PendStrict", f"Settings!$Q$4:$Q${3 + len(strict)}")
    for addr, lo, hi, msg in (("B5", "-12", "14", "UTC offset in hours, -12 to 14"),
                              ("B7", "1", "10", "Hours, 1 to 10 (must stay below the 12 h preparation check)"), ("B8", "1", "72", "Hours, 1 to 72"),
                              ("B9", "1", "48", "Hours, 1 to 48"), ("B10", "1", "60", "Days, 1 to 60")):
        dv = DataValidation(type="decimal", operator="between", formula1=lo, formula2=hi, allow_blank=False,
                            showErrorMessage=True, errorTitle="Invalid setting", error=msg)
        dv.add(addr)
        ws.add_data_validation(dv)
    dvb4 = DataValidation(type="decimal", operator="between", formula1="46023", formula2="46752", allow_blank=True,
                          showErrorMessage=True, errorTitle="Date-time required",
                          error="Enter a UTC date-time in 2026-2027, or leave blank for the live clock.")
    dvb4.add("B4")
    ws.add_data_validation(dvb4)

    ws["A11"] = "Station time zones (IANA tz database rules; DST windows cover the Oct-2026 programme)"
    ws["A11"].font = f(10, True, NAVY)
    header(ws, 12, ["Station", "IANA zone", "Std offset (h)", "DST offset (h)", "DST start (UTC)", "DST end (UTC)"], height=30)
    for c, w in zip("CDEF", (14, 14, 18, 18)):
        ws.column_dimensions[c].width = w
    for i, s in enumerate(data["stations"]):
        r = 13 + i
        vals = [s["code"], s["zone"], s["std"], s["dst"],
                dt(s["dst_start_utc"]) if s["dst_start_utc"] else None,
                dt(s["dst_end_utc"]) if s["dst_end_utc"] else None]
        for j, v in enumerate(vals, 1):
            c = ws.cell(r, j, v)
            c.font = f(9, color="0000FF")
            c.border = BORDER
            if j >= 5:
                c.number_format = DT
    assert len(data["stations"]) == 10
    for nm, col in (("TZ_STN", "A"), ("TZ_STD", "C"), ("TZ_DST", "D"), ("TZ_DSTS", "E"), ("TZ_DSTE", "F")):
        name(wb, nm, f"Settings!${col}$13:${col}$22")
    ws["A23"] = ("Source: IANA tz database. Europe/London BST ends Sun 25-Oct-2026 01:00 UTC; Australia/Adelaide ACDT "
                 "(UTC+10:30) starts Sun 4-Oct-2026 02:00 local (03-Oct 16:30 UTC). No DST at other stations.")
    ws["A23"].font = f(8, italic=True, color="595959")
    ws["A23"].alignment = WRAP
    ws.merge_cells("A23:F23")
    ws.row_dimensions[23].height = 30

    ws["H3"] = "Status list"
    ws["I3"] = "CA status list"
    for c in ("H3", "I3"):
        ws[c].font = f(9, True, NAVY)
    for i, v in enumerate(STATUS_LIST):
        ws.cell(4 + i, 8, v).font = f(9)
    for i, v in enumerate(CA_LIST):
        ws.cell(4 + i, 9, v).font = f(9)
    ws.column_dimensions["H"].width = 16
    ws.column_dimensions["I"].width = 14
    name(wb, "L_Status", "Settings!$H$4:$H$9")
    ws.print_area = "A1:F23"
    fit_pages(ws, "A", "F")
    ws.page_setup.fitToHeight = 1
    name(wb, "L_CA", "Settings!$I$4:$I$5")
    ws.sheet_view.showGridLines = False


# ------------------------------------------------------------------ Flights
FL_HEAD = ["Flight ID", "Itin", "Seq", "Date", "Day", "Flight No", "Dep", "Arr", "STD local", "STA local",
           "Class assessed", "Fleet (schedule)", "Aircraft ref type (STD UPLIFT INFO)", "Tail / Reg (input)",
           "Seat", "Transit", "Remark", "Sector key", "Region", "Service (ref)", "Ref uplift stns",
           "Meal uplift stn", "Round-trip catered from KUL: agenda candidate carrying flight", "Caterer (input)", "Flight PIC (input)",
           "Dep UTC offset (h)", "STD UTC", "Block time (h)", "T-7D due UTC", "T-24H due UTC",
           "T-12H prep due UTC", "Uplift due UTC (= STD)", "T-7D due local (meal uplift stn)",
           "T-24H due local", "T-12H prep due local", "T-7D %", "T-24H %", "T-12H prep %", "Uplift %",
           "Overall %", "Open required checks", "Overdue", "Open discrepancies", "Docs outstanding",
           "Clarifications open", "READINESS", "Invalid entries", "Checks completed", "Checks in scope",
           "Inbound carrying flight: KUL departure (UTC, input)", "KUL loading deadline for KUL-sourced items (UTC)",
           "Carrying flight earliest arrival at departure stn (UTC)", "First on-board confirmation (UTC)",
           "Wording flags to review"]
FL_W = [7, 5, 5, 10, 6, 9, 6, 6, 15, 15, 8, 9, 30, 12, 6, 8, 14, 13, 12, 22, 11, 9, 16, 14, 16,
        8, 15, 8, 15, 15, 15, 15, 15, 15, 15, 8, 8, 8, 8, 8, 9, 8, 9, 9, 9, 26, 8, 9, 9, 17, 17, 17, 15, 10]


def build_flights(wb, data, n_checks):
    ws = wb.create_sheet("Flights")
    title(ws, "Flight register",
          "Source: Skytrax Agenda 2026 slide (schedule) and STD UPLIFT INFORMATION.xlsx (uplift). Yellow = input. "
          "Times: STD/STA in local time of the departure/arrival station; due times computed from STD in UTC.",
          "G2:M2", 89)
    header(ws, 4, FL_HEAD, FL_W, height=54)
    last = 4 + n_checks
    rng = lambda col: f"Checks!${col}$5:${col}${last}"
    byid = {x["id"]: x for x in data["flights"]}
    for i, fl in enumerate(data["flights"]):
        r = 5 + i
        loaded = ""
        if fl["loaded_on"]:
            lf = byid[fl["loaded_on"]]
            loaded = lf["id"]
        static = [fl["id"], fl["itin"], fl["seq"], dt(fl["date"] + " 00:00"), fl["day"], fl["flt"], fl["dep"],
                  fl["arr"], dt(fl["std_local"]), dt(fl["sta_local"]), fl["cls"], fl["fleet"], fl["fleet_ref"],
                  None, fl["seat"], fl["transit"], fl["remark"], fl["sector"], fl["region"], fl["service"],
                  fl["ref_uplift_stns"], fl["meal_uplift_stn"], loaded, None, None]
        for j, v in enumerate(static, 1):
            ws.cell(r, j, v)
        ws.cell(r, 4).number_format = "dd-mmm-yy"
        ws.cell(r, 9).number_format = DT
        ws.cell(r, 10).number_format = DT
        # computed
        ws[f"Z{r}"] = "=" + off(f"G{r}", f"(I{r}-{stdoff(f'G{r}')}/24)")
        ws[f"AA{r}"] = f"=I{r}-Z{r}/24"
        arr_u0 = f"(J{r}-{stdoff('H' + str(r))}/24)"
        ws[f"AB{r}"] = f"=(J{r}-({off('H' + str(r), arr_u0)})/24-AA{r})*24"
        ws[f"AB{r}"].number_format = "0.00"
        ws[f"AC{r}"] = f"=AA{r}-7"
        if fl["round_trip"]:
            ws[f"AD{r}"] = f"=MIN(AA{r}-1,AY{r})"
            ws[f"AE{r}"] = f"=MIN(AA{r}-0.5,AY{r})"
        else:
            ws[f"AD{r}"] = f"=AA{r}-1"
            ws[f"AE{r}"] = f"=AA{r}-0.5"
        if fl["loaded_on"]:
            ws[f"W{r}"].comment = Comment(
                "Round-trip catered: reference uplift stn for this sector is KUL only, so return catering is loaded at KUL "
                f"on {byid[fl['loaded_on']]['flt']}. Preparation checks must finish before that loading window. "
                "Enter the actual carrying flight's KUL departure (UTC) in column AX; the deadlines follow it.", "MAGCS")
        ws[f"AF{r}"] = f"=AA{r}"
        for src, dst in (("AC", "AG"), ("AD", "AH"), ("AE", "AI")):
            ws[f"{dst}{r}"] = f"={src}{r}+({off(f'V{r}', f'{src}{r}')})/24"
        for col in ("AA", "AC", "AD", "AE", "AF", "AG", "AH", "AI"):
            ws[f"{col}{r}"].number_format = DT

        def pct(extra=""):
            return (f"=IFERROR(SUMIFS({rng('AJ')},{rng('B')},$A{r}{extra})/SUMIFS({rng('AI')},{rng('B')},$A{r}{extra}),\"n/a\")")
        ws[f"AJ{r}"] = pct(f",{rng('G')},\"T-7D\"")
        ws[f"AK{r}"] = pct(f",{rng('G')},\"T-24H\"")
        ws[f"AL{r}"] = pct(f",{rng('G')},\"T-12H PREP\"")
        ws[f"AM{r}"] = pct(f",{rng('G')},\"UPLIFT\"")
        ws[f"AN{r}"] = pct()
        for col in ("AJ", "AK", "AL", "AM", "AN"):
            ws[f"{col}{r}"].number_format = "0%"
            ws[f"{col}{r}"].alignment = Alignment(horizontal="right")
        ws[f"AO{r}"] = f"=AW{r}-AV{r}"
        ws[f"AP{r}"] = f"=SUMIFS({rng('AK')},{rng('B')},$A{r})"
        ws[f"AQ{r}"] = f"=SUMIFS({rng('AL')},{rng('B')},$A{r})"
        ws[f"AR{r}"] = f"=Documents!P{r}"
        tails = "{" + ",".join(f'\"9MMA{c}\"' for c in "BCDEFGH") + "}"
        tail_open = (f"IF(AND(TRIM(L{r})=\"A350\",ISNA(MATCH(UPPER(SUBSTITUTE(SUBSTITUTE(TRIM(N{r}),\"-\",\"\"),\" \",\"\")),"
                     f"{tails},0))),1,0)")  # A350 quantities / cart positions depend on the tail (9M-MAB..MAH)
        ws[f"AS{r}"] = f"=SUMIFS({rng('AN')},{rng('B')},$A{r})+{tail_open}"
        prep_done = (f"SUMIFS({rng('AJ')},{rng('B')},$A{r},{rng('H')},\"Preparation\")="
                     f"SUMIFS({rng('AI')},{rng('B')},$A{r},{rng('H')},\"Preparation\")")
        ws[f"AT{r}"] = (f"=IF(AND(AW{r}>0,AO{r}=0,AQ{r}=0,AR{r}=0,AS{r}=0,AU{r}=0),\"READY\","
                        f"IF(AP{r}>0,\"NOT READY {ND} OVERDUE\",IF(AQ{r}>0,\"NOT READY {ND} DISCREPANCY\","
                        f"IF(AU{r}>0,\"NOT READY {ND} INVALID ENTRY\",IF(AR{r}>0,\"NOT READY {ND} DOCUMENTS\","
                        f"IF(AS{r}>0,\"NOT READY {ND} CLARIFICATION\",IF(AND(AV{r}>0,{prep_done}),"
                        f"\"PREP DONE {ND} AWAITING UPLIFT\",IF(AV{r}=0,\"NOT STARTED\",\"IN PROGRESS\"))))))))")
        ws[f"AU{r}"] = f"=SUMIFS({rng('AM')},{rng('B')},$A{r})"
        ws[f"AV{r}"] = f"=SUMIFS({rng('AJ')},{rng('B')},$A{r})"
        ws[f"AW{r}"] = f"=SUMIFS({rng('AI')},{rng('B')},$A{r})"
        bounds = next((ck["due_bound"] for ck in data["checks"] if ck["flight_id"] == fl["id"] and ck.get("due_rule") == "CARRY"), None)
        if bounds:
            lo = (datetime.strptime(fl["std_utc"], "%Y-%m-%d %H:%M") - timedelta(days=3)).strftime("%Y-%m-%d %H:%M")
            prior = [b for b in bounds if lo <= byid[b]["std_utc"] < fl["std_utc"]]  # same 3-day rule as Checks
            if prior:
                default = f"$AA${5 + int(prior[-1][1:]) - 1}"
            else:
                mb = ",".join(f"$AB${5 + int(b[1:]) - 1}" for b in bounds)
                default = f"AA{r}-MIN({mb})/24"
            ws[f"AY{r}"] = f"=ROUND((IF(ISNUMBER(AX{r}),AX{r},{default})-UpliftWindowH/24)*1440,0)/1440"
            mb2 = ",".join(f"$AB${5 + int(b[1:]) - 1}" for b in bounds)
            az = f"ROUND((AY{r}+UpliftWindowH/24+MIN({mb2})/24)*1440,0)/1440"
            # without an agenda candidate the 'latest possible' carrier only caps preparation dues; it is not used to
            # reject on-board checks (the carrying-flight line itself must still be resolved before READY)
            ws[f"AZ{r}"] = f"={az}" if prior else f"=IF(ISNUMBER(AX{r}),{az},\"\")"
            ws[f"AX{r}"].comment = Comment(
                "Enter the UTC departure from KUL of the flight that carries this leg's KUL-sourced items. Until entered, the "
                "loading deadline assumes the latest possible departure (STD minus shortest agenda block).", "MAGCS")
        else:
            ws[f"AX{r}"] = "n/a"
            ws[f"AY{r}"] = "n/a"
            ws[f"AZ{r}"] = "n/a"
        start = f"IF(ISNUMBER($AZ{r}),MAX($AA{r}-UpliftWindowH/24,$AZ{r}),$AA{r}-UpliftWindowH/24)"
        mins = [f"_xlfn.MINIFS({rng('AG')},{rng('B')},$A{r},{rng('H')},\"Physical uplift\",{rng('P')},$G{r},"
                f"{rng('U')},\"{st}\",{rng('AG')},\">=\"&{start},{rng('AG')},\"<=\"&$AA{r})" for st in ("Pass", "Pass after CA")]
        first = "MIN(" + ",".join(f"IF({m}>0,{m},9E+99)" for m in mins) + ")"
        # first valid on-board confirmation: N/A, blank-status and out-of-window rows never count as loading
        ws[f"BA{r}"] = f"=IFERROR(IF({first}>1E+99,\"\",{first}),\"\")"
        ws[f"BB{r}"] = f"=SUMIFS({rng('BG')},{rng('B')},$A{r})"
        ws[f"AZ{r}"].number_format = DT
        ws[f"BA{r}"].number_format = DT
        ws[f"AX{r}"].number_format = DT
        ws[f"AY{r}"].number_format = DT
        for j in range(1, len(FL_HEAD) + 1):
            c = ws.cell(r, j)
            c.font = f(9)
            c.border = BORDER
            if c.alignment.horizontal is None:
                c.alignment = Alignment(vertical="top", wrap_text=True)
            if j in (14, 24, 25) or (j == 50 and ws.cell(r, 50).value != "n/a"):
                c.fill = F_INPUT
            elif j >= 26:
                c.fill = F_CALC
        ws[f"L{r}"].fill = F_INPUT
        ws[f"L{r}"].comment = None
        ws[f"AT{r}"].font = f(9, True)
        ws.row_dimensions[r].height = 36
    ws["L4"].comment = Comment("Editable: if an equipment change occurs, update the fleet – the A339 IFE-menu rule follows this cell.", "MAGCS")
    rngF = "A5:AW26"
    ws.conditional_formatting.add("AT5:AT26", CellIsRule(operator="equal", formula=['"READY"'],
                                  fill=PatternFill("solid", fgColor="C6EFCE"), font=Font(name=FONT, color="006100", bold=True)))
    ws.conditional_formatting.add("AT5:AT26", FormulaRule(formula=['LEFT(AT5,9)="NOT READY"'],
                                  fill=PatternFill("solid", fgColor="FFC7CE"), font=Font(name=FONT, color="9C0006", bold=True)))
    ws.conditional_formatting.add("AT5:AT26", FormulaRule(formula=['LEFT(AT5,9)="PREP DONE"'],
                                  fill=PatternFill("solid", fgColor="FFEB9C")))
    ws.conditional_formatting.add("AP5:AS26", CellIsRule(operator="greaterThan", formula=["0"],
                                  font=Font(name=FONT, color="C00000", bold=True)))
    ws.freeze_panes = "G5"
    ws.auto_filter.ref = f"A4:BB26"
    dvx = DataValidation(type="decimal", operator="between", formula1="46204", formula2="46419", allow_blank=True,
                         showErrorMessage=True, error="Enter the carrying flight's KUL departure as a UTC date-time (Jul 2026 - Jan 2027).")
    dvx.add("AX5:AX26")
    dvl = DataValidation(type="list", formula1='"A350,A333,A339,B738MAX"', allow_blank=False, showErrorMessage=True,
                         errorTitle="Fleet", error="Choose the fleet from the list (schedule values).")
    dvl.add("L5:L26")
    ws.add_data_validation(dvl)
    ws.add_data_validation(dvx)
    ws.sheet_view.zoomScale = 85
    ws.print_options.gridLines = False
    fit_pages(ws, "A", "BB", title_cols_w=43)
    ws.print_title_rows = "4:4"
    ws.print_title_cols = "A:F"


# ------------------------------------------------------------------ Checks
def norm(ref):
    """Lower-case text with spaces, non-breaking spaces and punctuation removed (placeholder / identity tests)."""
    x = f"LOWER({ref}&\"\")"
    for ch in ('CHAR(10)', 'CHAR(13)', 'CHAR(9)', 'CHAR(160)', '" "', '"."', '"-"', '"/"', '"?"', '"_"', '","', '"["', '"]"', '"("', '")"', '"*"'):
        x = f"SUBSTITUTE({x},{ch},\"\")"
    return x


ROLE_WORDS = ["qa", "qc", "qaqc", "csm", "cic", "ccic", "ops", "pic", "self", "team", "staff", "crew", "cabin", "caterer",
              "catering", "pasb", "mcat", "sats", "magcs", "mas", "mh", "supervisor", "spvr", "supv", "sup", "manager",
              "mgr", "officer", "offr", "steward", "stewardess", "purser", "agent", "agents", "rep", "reps",
              "representative", "dept", "department", "shift", "night", "day", "morning", "evening", "ramp", "loading",
              "hygiene", "inflight", "services", "service", "quality", "assurance", "control", "controller", "chief",
              "head", "lead", "leader", "senior", "sr", "jr", "junior", "station", "duty", "inspector", "executive",
              "exec", "captain", "capt", "chef", "verified", "verifier", "checked", "by", "the", "of", "and", "on", "desk",
              "unit", "section", "office", "store", "stores", "kitchen", "production", "planning", "dispatch",
              "coordinator", "assistant", "asst", "admin", "security", "engineer", "tech", "technician", "flight",
              "mr", "mrs", "ms", "dr", "en", "puan", "cik", "kul", "lhr", "pen", "lgk", "cgk", "han", "hkg", "kix",
              "nrt", "adl", "in", "charge", "person", "incharge", "airline", "aircraft", "operations", "operation",
              "attendant", "attendants", "fa", "sfa", "ground", "handling", "handler", "loadmaster", "foreman",
              "hostess", "sous", "galley", "leading", "hand", "lae", "driver", "loader", "loaders", "dispatcher",
              "supervisors", "managers", "officers", "stewards", "chefs", "inspectors", "agents", "cook", "porter",
              "cleaner", "mechanic", "engineer", "crewmember", "member", "members", "cabincrew", "groundstaff"]
# N/A justification filler: what is left after removing these must still say something (8+ letters)
NA_FILLER = ["not", "applicable", "applies", "apply", "for", "this", "that", "the", "flight", "flights", "sector",
             "leg", "na", "n", "a", "required", "needed", "need", "is", "it", "on", "same", "as", "previous", "above",
             "below", "see", "refer", "line", "row", "item", "ditto", "per", "check", "checks", "nil", "none", "no",
             "relevant", "case", "here", "now", "today", "x", "xx", "xxx", "to", "of", "and", "be", "one", "other",
             "prev", "ref", "note", "noted", "ok"]


def name_letters(spaced_cell, lst="L_Role"):
    """Letters left in a spaced ' word word ' cell after removing every whole word found in a Settings list."""
    removed = (f"SUMPRODUCT((LEN({spaced_cell})-LEN(SUBSTITUTE({spaced_cell},\" \"&{lst}&\" \",\" \")))"
               f"/(LEN({lst})+1)*LEN({lst}))")
    digits = f"SUMPRODUCT(--ISNUMBER(FIND({{\"0\",\"1\",\"2\",\"3\",\"4\",\"5\",\"6\",\"7\",\"8\",\"9\"}},{spaced_cell})))"
    return f"IF({digits}>0,0,LEN(SUBSTITUTE({spaced_cell},\" \",\"\"))-{removed})"


def same_person(sp_a, sp_b, n_a, n_b):
    """Same name in another order ('Aisyah Rahman' / 'Rahman, Aisyah'): equal letters and every word of A in B."""
    words = [f"TRIM(MID(SUBSTITUTE(TRIM({sp_a}),\" \",REPT(\" \",60)),{k * 60 + 1},60))" for k in range(3)]
    found = ",".join(f"OR({w}=\"\",ISNUMBER(SEARCH(\" \"&{w}&\" \",{sp_b})))" for w in words)
    last = lambda sp: f"TRIM(RIGHT(SUBSTITUTE(TRIM({sp}),\" \",REPT(\" \",60)),60))"
    patr = f"SUMPRODUCT(--ISNUMBER(SEARCH({{\" bin \",\" binti \",\" bt \",\" bte \",\" a l \",\" a p \"}},{sp_a}&{sp_b})))>0"
    first = lambda sp: f"LEN(TRIM(LEFT(TRIM({sp}),FIND(\" \",TRIM({sp})&\" \"))))"
    initial = (f"AND(LEN({last(sp_a)})>=3,{last(sp_a)}={last(sp_b)},LEFT(TRIM({sp_a}),1)=LEFT(TRIM({sp_b}),1),"
               f"OR({first(sp_a)}=1,{first(sp_b)}=1),"
               f"NOT({patr}))")
    words_b = [f"TRIM(MID(SUBSTITUTE(TRIM({sp_b}),\" \",REPT(\" \",60)),{k * 60 + 1},60))" for k in range(3)]
    found_ba = ",".join(f"OR({w}=\"\",ISNUMBER(SEARCH(\" \"&{w}&\" \",{sp_a})))" for w in words_b)
    same_len = f"LEN({n_a})=LEN({n_b})"
    return (f"OR(AND(LEN({n_a})>=3,{found},NOT(AND({patr},{same_len}))),"
            f"AND(LEN({n_b})>=3,{found_ba},NOT(AND({patr},{same_len}))),{initial})")


def spaced(ref):
    """' word word ' form: lower case, punctuation turned into single spaces, padded (whole-word SEARCH)."""
    x = f"LOWER({ref}&\"\")"
    for ch in ('CHAR(10)', 'CHAR(13)', 'CHAR(9)', 'CHAR(160)', '"."', '","', '"-"', '"/"', '"("', '")"', '":"', '";"', '"_"', '"–"', '"?"', '"!"',
               '"*"', '"["', '"]"', '"\'"', '"#"', '"+"', '"&"', '"\\"'):
        x = f"SUBSTITUTE({x},{ch},\" \")"
    return f"\" \"&TRIM({x})&\" \""


def occurrences(listname, s):
    """Total whole-word occurrences in spaced text s of every entry of a list (occurrence counting)."""
    return (f"SUMPRODUCT((LEN({s})-LEN(SUBSTITUTE({s},\" \"&{listname}&\" \",\" \")))/(LEN({listname})+1))")


def pending(s, lst="L_Pending"):
    return f"{occurrences(lst, s)}>{occurrences('L_PendNeg', s)}"


def digitmap(ref):
    x = f"LOWER({ref}&\"\")"
    for d in "123456789":
        x = f"SUBSTITUTE({x},\"{d}\",\"0\")"
    return x


def alphamap(ref):
    x = ref
    for ch in "bcdefghijklmnopqrstuvwxyz":
        x = f"SUBSTITUTE({x},\"{ch}\",\"a\")"
    return x


REF_WORDS = ["ref", "no", "nos", "#", "number", "numbers", "photo", "photos", "picture", "pictures", "pic", "image",
             "seal", "seals", "receipt", "note", "dn", "invoice", "po", "voucher", "checklist",
             "scan", "manifest", "binder", "folder", "file", "email", "memo", "form", "sheet", "log", "doc", "report",
             "img", "ack", "batch", "lot", "job", "record", "fax", "letter", "ticket", "case", "order", "id"]
# words that precede numbers that are NOT record references (times, dates, places) – masked before the ID test
MASK_WORDS_1 = ["at", "by", "on", "to", "from", "until", "till", "before", "after", "around", "approx", "about",
                "since", "hrs", "hr", "hours", "time", "jan", "feb", "mar", "apr", "may", "jun", "jul", "aug", "sep",
                "sept", "oct"]
MASK_WORDS_2 = ["nov", "dec", "mon", "tue", "wed", "thu", "fri", "sat", "sun", "gate", "bay", "stand", "position",
                "galley", "flight", "flt", "day", "date", "dated", "the", "and", "of", "in", "for", "with", "pax",
                "etd", "eta", "std", "sta", "atd", "ata", "lt", "utc", "local", "ok", "done", "seen", "chk", "chkd",
                "checked", "rcvd", "recd", "dep", "arr", "qty", "all", "ttl", "total", "via", "per", "call", "inbound",
                "outbound", "approx"]
MONTHS = ["jan", "feb", "mar", "apr", "may", "jun", "jul", "aug", "sep", "oct", "nov", "dec"]


def masked(spaced_cell, flight_no_cell, stage):
    """Mask time/date/place/ops words and any carrier-coded flight number ('mh0727', 'mh 727') with '~'.
    Two stages (separate helper cells) keep each formula well inside Excel's 64-level nesting limit."""
    if stage == 1:
        x = f"SUBSTITUTE(SUBSTITUTE({spaced_cell},\" \"&LOWER({flight_no_cell})&\" \",\" ~ \"),\" mh\",\" ~\")"
        words = MASK_WORDS_1
    elif stage == 2:
        x = spaced_cell
        for m in MONTHS:  # dates written without spaces, e.g. 09oct2026
            x = f"SUBSTITUTE({x},\"{m}\",\"~\")"
        words = MASK_WORDS_2[:25]
    else:
        x = spaced_cell
        words = MASK_WORDS_2[25:]
    for w in words:
        x = f"SUBSTITUTE({x},\" {w} \",\" ~ \")"
    return x


def ref_ok(raw, dmap, amap, batch=False):
    """Traceable reference: 2+ letters, optional separator, 3+ digits (SF-2210, IMG_2231, DN 88213), a reference
    keyword followed by a number (seal no 88213, email 4471), a file name, link or path. Times, dates, flight
    numbers and galley positions are masked out first, so 'Checked at 1400' or 'Checked 09-Oct-26' do not count."""
    alnum = '{\"aa000\",\" aa 000\",\" aaa 000\",\" aaaa 000\",\" a 0000\",\" a000 000\",\" aa 00 000\",\" aaa 00 000\"}'
    kw = "{" + ",".join(f'\" {w} 00\"' for w in REF_WORDS) + "," + ",".join(f'\" {w}00\"' for w in REF_WORDS) + "}"
    ext = ("{\"http\",\"www.\",\"\\\",\".pdf\",\".jpg\",\".jpeg\",\".png\",\".heic\",\".xls\",\".doc\","
           "\".msg\",\".eml\"}")
    extra = f",ISNUMBER(SEARCH(\"0000\",{dmap}))" if batch else ""
    # a named document store plus a folder/item number (e.g. "SharePoint F11/UPL-01")
    extra += (f",AND(SUMPRODUCT(--ISNUMBER(SEARCH({{\"sharepoint\",\"onedrive\",\"teams\",\"dms\",\"q-pulse\",\"qpulse\"}},{raw})))>0,"
              f"ISNUMBER(SEARCH(\"0\",{dmap})))")
    return (f"OR(SUMPRODUCT(--ISNUMBER(SEARCH({alnum},{amap})))>0,SUMPRODUCT(--ISNUMBER(SEARCH({kw},{dmap})))>0,"
            f"SUMPRODUCT(--ISNUMBER(SEARCH({ext},{raw})))>0{extra})")


OUTCOMES = ["Confirmed – applies / carried as listed", "Confirmed – not carried on this sector",
            "Confirmed – not applicable to this aircraft / tail", "Refreshment service – no printed menu card"]
EVIDENCE_TYPES = ["Form / checklist", "Sensory panel sheet", "Load / uplift sheet", "Photo", "Email", "Seal no.",
                  "Delivery note / receipt", "Memo / letter", "Document (GLD / menu / ISOP)", "System record",
                  "Link / file path"]
# real links / paths / file names only (a bare "/" is NOT a link: 9/10, A/C, QA/QC)
LINK_EXT = ("{\"http\",\"www.\",\"\\\\\",\":\\\",\".pdf\",\".jpg\",\".jpeg\",\".png\",\".heic\",\".xls\",\".doc\","
            "\".msg\",\".eml\",\".ppt\",\".txt\",\".zip\"}")
JOINED = '{\"a0\",\"a-0\",\"a_0\",\"a/0\",\"a.0\",\"a#0\",\"0a\",\"0-a\",\"0_a\",\"0/a\"}'
BAD_SHAPES = ('{\"0.00.00\",\"0.0.00\",\"00.0.00\",\"000a\",\"00a\",\"0000aa\",\"000aa\",\"00aa\",\"0aa\",\"0000aaa\",\"000aaa\",\"00a00\",'
              '\"0a00\",\"0000-0000\",\"00:00\",\"a00\",\"a0\",\"00.00.0000\",\"00/00/0000\",\"00-00-0000\",'
              '\"0000-00-00\",\"0000/00/00\",\"0000.00.00\",\"00.00.00\",\"00/00/00\",\"00-00-00\",\"0.00.0000\",'
              '\"0/00/0000\",\"0-00-0000\"}')
# month / time-zone / revision words: letters that do not make a token an ID (9OCT26, 1415MYT)
STRIP = ["revision", "version", "jan", "feb", "mar", "apr", "may", "jun", "jul", "aug", "sep", "oct", "nov", "dec",
         "utc", "gmt", "myt", "bst", "rev", "ver"]
# leading words that mark a token as something other than a record ID (flight, revision, aircraft, seat, time)
NOT_ID_PREFIX = ["mh", "rev", "revision", "ver", "version", "v", "r", "q", "fy", "seat", "tray", "row", "gate", "bay",
                 "stand", "page", "pg", "item", "line", "cart", "galley", "flight", "flt", "table", "box", "bag",
                 "oven", "etd", "eta", "std", "sta", "atd", "ata", "qty", "pax", "x", "time", "tel", "ph",
                 "mob", "day", "wk", "week", "utc", "gmt", "myt", "lt"]
DUMMY_LETTERS = ["x", "xx", "xxx", "xxxx", "xxxxx", "abc", "abcd", "nil", "na", "nan", "test", "tbc", "tba", "tbd",
                 "dummy", "sample", "example", "eg", "zz", "zzz", "none", "null", "xyz", "qwe", "asd", "id", "ref"]
DUMMY_DIGITS = ["12345", "123456", "1234567", "12345678", "01234", "012345", "98765",
                "987654", "54321"]
SEPS = ["-", "_", "/", ".", "#", "\\\\", ":", " ", "(", ")"]


def _arr(items):
    return "{" + ",".join('\"' + w + '\"' for w in items) + "}"


def stripped(ref):
    """Digit-mapped text with month / time-zone / revision words removed (9OCT26, 1415MYT, Rev3 lose their letters)."""
    x = ref
    for w in STRIP:
        x = f"SUBSTITUTE({x},\"{w}\",\"\")"
    return x


def digits_only(ref):
    """The digits of a cell, in order (letters and separators removed)."""
    x = f"LOWER({ref}&\"\")"
    for ch in "abcdefghijklmnopqrstuvwxyz":
        x = f"SUBSTITUTE({x},\"{ch}\",\"\")"
    for ch in SEPS:
        x = f"SUBSTITUTE({x},\"{ch}\",\"\")"
    return x


def letters_only(dmap):
    """The letters of a digit-mapped cell (digits are '0' there), separators removed."""
    x = f"SUBSTITUTE({dmap},\"0\",\"\")"
    for ch in SEPS:
        x = f"SUBSTITUTE({x},\"{ch}\",\"\")"
    return x


def alpha_prefix(dmap):
    """Letters before the first digit (separators removed): 'rev', 'mh', 'seat', 'sf' ..."""
    x = f"LEFT({dmap},FIND(\"0\",{dmap}&\"0\")-1)"
    for ch in SEPS:
        x = f"SUBSTITUTE({x},\"{ch}\",\"\")"
    return x


def _count(amap, ch):
    return f"(LEN({amap})-LEN(SUBSTITUTE({amap},\"{ch}\",\"\")))"


def _date6(dig):
    """Six digits that read as a 2025-27 date: yymmdd (261008) or ddmmyy (081026)."""
    mm = f"IFERROR(--MID({dig},3,2),0)"
    yrs = '{\"25\",\"26\",\"27\"}'
    return (f"AND(LEN({dig})=6,{mm}>=1,{mm}<=12,OR(ISNUMBER(MATCH(LEFT({dig},2),{yrs},0)),"
            f"ISNUMBER(MATCH(RIGHT({dig},2),{yrs},0))))")


def _dummy_or_date(dig, let):
    """Placeholder numbers (00000, 11111, 12345, SF-00), placeholder letters (XX, ABC, NA) and dates / timestamps."""
    return (f"OR(AND({dig}<>\"\",SUBSTITUTE({dig},\"0\",\"\")=\"\"),"
            f"AND({let}=\"\",LEN({dig})>=4,{dig}=REPT(LEFT({dig}),LEN({dig}))),"
            f"ISNUMBER(MATCH({dig},{_arr(DUMMY_DIGITS)},0)),ISNUMBER(MATCH({let},{_arr(DUMMY_LETTERS)},0)),"
            f"AND(OR({let}=\"\",{let}=\"t\"),LEN({dig})>=8,OR(LEFT({dig},3)=\"202\",AND(LEN({dig})=8,"
            f"ISNUMBER(MATCH(RIGHT({dig},4),{{\"2025\",\"2026\",\"2027\"}},0))))),"
            f"AND({let}=\"\",{_date6(dig)}))")


def id_shape_ok(raw, full, strp, prefix, dig, let, typ, flight):
    """Evidence ID is ONE token that is structurally a record ID: letters (other than month / zone / revision words)
    joined to 2+ digits (SF-2210, IMG_2231, LS/OCT/0118, DN88213) or a 5+ digit number (seal 1045521).
    Links, paths and file names only for type 'Link / file path'. Rejected: spaces or several IDs, flight numbers,
    dates, times, revisions, aircraft types, seat / tray numbers and placeholder values (00000, XX-0000, 12345)."""
    t = f"TRIM({raw})"
    link = f"SUMPRODUCT(--ISNUMBER(SEARCH({LINK_EXT},{raw})))>0"
    multi = (f"OR(SUMPRODUCT(--ISNUMBER(FIND({{\",\",\";\",\"&\",\"+\"}},{raw})))>0,"
             f"SUMPRODUCT(--ISNUMBER(SEARCH({{\"0/aa\",\"0\\\\aa\",\"-0000/0000\",\"_0000/0000\",\"-000/000\",\"-00000/00000\"}},{full})))>0)")
    aircraft = f"ISNUMBER(MATCH(LEFT(LOWER({t}),4),{_arr(['a350', 'a359', 'a330', 'a332', 'a333', 'a339', 'b737', 'b738', 'b7m8', 'a320', 'a321'])},0))"
    slashes = f"(LEN({raw})-LEN(SUBSTITUTE({raw},\"/\",\"\")))"
    ext = ("{\".pdf\",\".jpg\",\".jpeg\",\".png\",\".heic\",\".xls\",\".doc\",\".msg\",\".eml\",\".ppt\","
           "\".txt\",\".zip\",\".mp4\",\".mov\"}")
    personal = '{"\\users\\","\\desktop","\\downloads","\\documents\\"}'
    link_ok = (f"AND(LEN({t})>=8,SUMPRODUCT(--ISNUMBER(SEARCH({personal},{raw})))=0,"
               f"OR(SUMPRODUCT(--ISNUMBER(SEARCH({ext},{raw})))>0,AND(LEFT(LOWER({t}),4)=\"http\",{slashes}>=4)))")
    flight_like = f"LOWER({t})=LOWER({flight})"
    bad_shape = (f"OR(ISNUMBER(MATCH(LEFT({full},255),{BAD_SHAPES},0)),"
                 f"ISNUMBER(MATCH(LEFT({strp},255),{BAD_SHAPES},0)))")
    joined = f"SUMPRODUCT(--ISNUMBER(SEARCH({JOINED},{full})))>0"
    token = (f"AND(ISERROR(FIND(\" \",{t})),NOT({multi}),NOT({link}),NOT({flight_like}),NOT({aircraft}),"
             f"ISNA(MATCH({prefix},{_arr(NOT_ID_PREFIX)},0)),NOT({bad_shape}),OR(NOT({_dummy_or_date(dig, let)}),AND(TRIM({typ})=\"Photo\",ISNUMBER(FIND(\"_\",{t})),LEN({dig})>=14)),"
             f"OR(AND({_count(strp, 'a')}>=1,{_count(full, '0')}>=2,{joined}),"
             f"AND({_count(full, 'a')}=0,{_count(full, '0')}>=5)))")
    return f"IF(TRIM({typ})=\"Link / file path\",{link_ok},{token})"


def batch_shape_ok(raw, full, strp, dig, flight):
    """Batch ID: letters (not month words) joined to digits (PASB-261008-BC-017, LOT-2410-B7) or a 6+ digit
    production code (261008-017); at least 3 digits; not a flight number, date or placeholder."""
    return (f"AND(ISERROR(SEARCH(\" mh\",\" \"&LOWER({raw}))),ISNA(MATCH(LEFT(TRIM({full}),255),{BAD_SHAPES},0)),"
            f"ISNA(MATCH(LEFT(TRIM({strp}),255),{BAD_SHAPES},0)),LEN({dig})>=3,NOT({_date6(dig)}),"
            f"SUMPRODUCT(--ISNUMBER(FIND({{\",\",\";\",\"&\",\"+\"}},{raw})))=0,"
            f"NOT(OR(SUBSTITUTE({dig},\"0\",\"\")=\"\",AND(LEN({dig})>=4,{dig}=REPT(LEFT({dig}),LEN({dig}))),"
            f"ISNUMBER(MATCH({dig},{_arr(DUMMY_DIGITS)},0)),AND(LEN({dig})=8,LEFT({dig},3)=\"202\"))),"
            f"OR(AND({_count(strp, 'a')}>=1,SUMPRODUCT(--ISNUMBER(SEARCH({JOINED},{full})))>0),"
            f"ISNUMBER(SEARCH(\"000000\",{full}))))")


def has_ref(ref):
    """Traceable reference: a digit, a link, a path or a file name."""
    return (f"OR(SUMPRODUCT(--ISNUMBER(FIND({{\"0\",\"1\",\"2\",\"3\",\"4\",\"5\",\"6\",\"7\",\"8\",\"9\"}},{ref})))>0,"
            f"SUMPRODUCT(--ISNUMBER(SEARCH({{\"http\",\"www.\",\"\\\",\".pdf\",\".jpg\",\".jpeg\",\".png\",\".heic\","
            f"\".xls\",\".doc\",\".msg\",\".eml\"}},{ref})))>0)")


def has(listname, spaced_cell):
    return f"SUMPRODUCT(--ISNUMBER(SEARCH(\" \"&{listname}&\" \",{spaced_cell})))>0"


def due_formula(ck, cp, fr, code_due):
    rule = ck.get("due_rule", "")
    if rule == "CARRY" and cp == "UPLIFT":
        return f"=Flights!$AY${fr}+UpliftWindowH/24"
    if rule == "CARRY":
        return f"=MIN(Flights!${code_due[cp]}${fr},Flights!$AY${fr})"
    if rule == "LOADFLT":
        raise ValueError("LOADFLT rule retired")
    return f"=Flights!${code_due[cp]}${fr}"


CK_HEAD = ["Check ID", "Flight ID", "Flight No", "Date", "Sector", "Class", "Checkpoint", "Check type", "Category",
           "Check item", "Requirement / expected (reference)", "Source reference", "Item uplift stn",
           "Applicability", "Rule note / clarification question", "Check station", "Due (UTC)",
           "Due (local @ check stn)", "Evidence type",
           "PIC", "Status", "Result / assessment", "Batch ID (T-24H)", "Expected qty", "Actual qty",
           "Evidence ID / reference", "Corrective action", "CA status", "Outcome (clarification / N/A)",
           "Completion time (local @ check stn)", "Verifier",
           "Qty variance", "Completion (UTC)", "RECORD STATE", "In scope", "Complete", "Overdue",
           "Open discrepancy", "Invalid", "Clarification open", "Overdue seq", "Discrepancy seq",
           "Req batch", "Req qty", "Req doc", "N/A permitted", "n PIC", "n Result", "n Evidence", "n Verifier",
           "n Batch", "n N/A just.", "n CA", "CA / qty messages", "s Result", "s Evidence", "s Batch", "s N/A just.", "Wording flag",
           "Pending: result", "Pending: evidence", "Pending: batch", "Pending: N/A just.", "Evidence has ref",
           "d Evidence", "a Evidence", "d Batch", "a Batch", "Batch has ID", "(spare)", "(spare)",
           "(spare)", "(spare)", "Not on board", "(spare)", "(spare)", "Earliest valid (UTC)", "Evidence ID shape ok"]
CK_W = [14, 6, 9, 10, 9, 6, 11, 11, 12, 38, 38, 30, 9, 13, 40, 8, 15, 15, 15,
        14, 13, 40, 16, 9, 9, 22, 40, 9, 26, 17, 18,
        9, 15, 34, 6, 7, 7, 9, 7, 9, 8, 9, 6, 6, 6, 8, 8, 8, 8, 8, 8, 8, 8, 20, 12, 12, 12, 12, 8, 8, 8, 8, 8, 8, 10, 10, 10, 10, 8, 10, 10, 10, 10, 8, 10, 10, 15, 8]


def build_checks(wb, data):
    ws = wb.create_sheet("Checks")
    title(ws, "Check register",
          "Every readiness record for every flight (single source of truth). Enter results only in the yellow columns S–AE. Blank never counts as complete. Pass requires completion time, "
          "evidence and verifier (plus batch ID / quantities where required). N/A requires justification and verifier.",
          "D2:J2", 97)
    header(ws, 4, CK_HEAD, CK_W, height=54)
    fidx = {x["id"]: i for i, x in enumerate(data["flights"])}
    code_due = {"T-7D": "AC", "T-24H": "AD", "T-12H PREP": "AE", "UPLIFT": "AF"}
    early = {"T-7D": "T7EarlyD*24", "T-24H": "T24EarlyH", "T-12H PREP": "T12EarlyH", "UPLIFT": "UpliftWindowH"}
    n = len(data["checks"])
    last = 4 + n
    rowmap = {ck["check_id"]: 5 + i for i, ck in enumerate(data["checks"])}
    sqmap = {}
    for q in station_questions(data):
        sqmap[q["prep"]] = q
        sqmap[q["upl"]] = q
    for i, ck in enumerate(data["checks"]):
        r = 5 + i
        fr = 5 + fidx[ck["flight_id"]]
        cp = ck["checkpoint"]
        applic = ck["applic"]
        if applic == "RULE:A339":
            applic = f"=IF(TRIM(Flights!$L${fr})=\"A339\",\"Required\",\"N/A {ND} rule\")"
        elif applic == "RULE:LOADFLT":
            applic = f"=IF(TRIM(Flights!$W${fr})=\"\",\"N/A {ND} rule\",\"Required\")"
        exp_qty = None
        if "per tail" in ck["exp_qty"]:
            exp_qty = f"=IF(TRIM(Flights!$N${fr})=\"\",\"\",IF(ISNUMBER(SEARCH(\"MAH\",Flights!$N${fr})),280,260))"
        elif ck["exp_qty"]:
            try:
                exp_qty = int(ck["exp_qty"])
            except ValueError:
                exp_qty = None
        vals = {
            "A": ck["check_id"], "B": ck["flight_id"], "C": f"=Flights!$F${fr}", "D": f"=Flights!$D${fr}",
            "E": f"=Flights!$G${fr}&\"-\"&Flights!$H${fr}", "F": f"=Flights!$K${fr}", "G": cp,
            "H": ck["check_type"], "I": ck["category"], "J": ck["item"], "K": ck["expected"], "L": ck["src"],
            "M": ck["uplift_stn"], "N": applic, "O": ck["note"], "P": ck["station"],
            "Q": "=ROUND((" + due_formula(ck, cp, fr, code_due)[1:] + ")*1440,0)/1440",
            "R": f"=Q{r}+({off(f'P{r}', f'Q{r}')})/24",
            "BY": (f"=IF(ISNUMBER(Flights!$AZ${fr}),MAX(Q{r}-UpliftWindowH/24,Flights!$AZ${fr}),"
                  f"Q{r}-UpliftWindowH/24)" if cp == "UPLIFT" and ck.get("due_rule") != "CARRY"
                  else "" if (cp == "T-7D" and ck["applic"] == "Clarification required")
                  else f"=Q{r}-{early[cp]}/24"),
            "X": exp_qty,
            "AF": f"=IF(AND(ISNUMBER(X{r}),ISNUMBER(Y{r})),Y{r}-X{r},\"\")",
            "AG": f"=IF(ISNUMBER(AD{r}),AD{r}-({off('P' + str(r), '(AD' + str(r) + '-' + stdoff('P' + str(r)) + '/24)')})/24,\"\")",
            "AQ": ck["req_batch"], "AR": ck["req_qty"], "AS": ck["req_doc"] if ck["req_doc"] else 0,
        }
        sq = sqmap.get(ck["check_id"])
        if sq:  # uplift station confirmed on Settings (station confirmations) – the line follows the answer
            sr = sq["row"]
            conf, carr, arr = f"TRIM(Settings!$F${sr})", f"Settings!$G${sr}", f"Settings!$I${sr}"
            moved = f"AND({conf}<>\"\",{conf}<>\"{sq['dep']}\")"
            if ck["check_type"] == "Preparation":
                vals["P"] = f"=IF({conf}=\"\",\"{sq['default']}\",{conf})"
                base = vals["Q"][1:]
                vals["Q"] = f"=IF(AND({moved},ISNUMBER({carr})),MIN({base},{carr}-UpliftWindowH/24),{base})"
            else:
                vals["BY"] = f"=MAX({vals['BY'][1:]},IF(AND({moved},ISNUMBER({arr})),{arr},0))"
        # doc linkage: T7 GLD row = req_doc 1, menu checklist row = 2
        if ck["req_doc"]:
            vals["AS"] = 1 if "GLD" in ck["item"] else 2
        if ck.get("req_carry"):
            vals["AS"] = 3
        doc_chk = (f"IF(AND(AS{r}=1,Documents!$J${fr}<>\"ON FILE\"),\"INVALID {ND} GLD not on file (Documents sheet)\","
                   f"IF(AND(AS{r}=1,ISNUMBER(Documents!$H${fr}),Documents!$H${fr}>AD{r}),\"INVALID {ND} GLD revised after this check (Documents rev date): re-check against the current revision\","
                   f"IF(AND(AS{r}=2,ISNUMBER(Documents!$M${fr}),Documents!$M${fr}>AD{r}),\"INVALID {ND} menu checklist revised after this check (Documents rev date): re-check against the current revision\","
                   f"IF(AND(AS{r}=2,Documents!$O${fr}<>\"ON FILE\"),\"INVALID {ND} menu checklist not on file (Documents sheet)\","
                   f"IF(AND(AS{r}=3,NOT(ISNUMBER(Flights!$AX${fr}))),\"INVALID {ND} enter carrying flight KUL departure (Flights AX)\","
                   f"IF(AND(AS{r}=3,OR(Flights!$AZ${fr}>Flights!$AA${fr},Flights!$AX${fr}<Flights!$AA${fr}-3)),"
                   f"\"INVALID {ND} carrying flight (Flights AX) cannot arrive before this leg's STD, or is over 3 days early\",")
        # msg = qty-variance / CA checks joined with &; "" when OK. IF(msg<>"", msg, rest) keeps Excel-2007 nesting.
        U = f"TRIM(U{r})"

        helper = {"T": "AU", "V": "AV", "Z": "AW", "AE": "AX", "W": "AY", "AC": "AZ", "AA": "BA"}
        for src, hcol in helper.items():
            vals[hcol] = "=" + norm(f"{src}{r}")
        vals["CC"] = "=" + spaced(f"T{r}")
        vals["CD"] = "=" + spaced(f"AE{r}")
        vals["BR"] = f"=IFERROR({name_letters(f'CC{r}')},0)"  # PIC: letters left after role / title words
        vals["BS"] = f"=IFERROR({name_letters(f'CD{r}')},0)"  # verifier: letters left after role / title words
        vals["CE"] = f"=IFERROR(IF({same_person(f'CC{r}', f'CD{r}', f'AU{r}', f'AX{r}')},1,0),0)"
        sp = {"V": "BC", "Z": "BD", "W": "BE", "AC": "BF"}
        for src, hcol in sp.items():
            vals[hcol] = "=" + spaced(f"{src}{r}")
        prep_row = ck["check_type"] == "Preparation"
        for src, sc, hcol, lst in (("V", "BC", "BH", "L_Pending"), ("Z", "BD", "BI", "L_PendStrict"),
                                   ("W", "BE", "BJ", "L_Pending"), ("AC", "BF", "BK", "L_PendStrict")):
            extra = f"+{occurrences('L_PendPrep', f'{sc}{r}')}" if (prep_row and src in ("V", "W")) else ""
            vals[hcol] = (f"=IFERROR(IF({occurrences(lst, f'{sc}{r}')}>{occurrences('L_PendNeg', f'{sc}{r}')}{extra},1,0),1)")
        # on-board rows: result or evidence saying the item is not on board blocks the confirmation
        vals["BV"] = (f"=IFERROR(IF(OR({occurrences('L_NotLoaded', f'BC{r}')}>{occurrences('L_NotLoadedNeg', f'BC{r}')},"
                      f"{occurrences('L_NotLoaded', f'BD{r}')}>{occurrences('L_NotLoadedNeg', f'BD{r}')}),1,0),0)"
                      if not prep_row else "=0")
        # structural evidence ID: shape checked against the chosen evidence type (column S)
        vals["BM"] = "=" + digitmap(f"Z{r}")
        vals["BN"] = "=" + alphamap(f"BM{r}")                 # full shape: letters a, digits 0
        vals["BT"] = "=" + alphamap(stripped(f"BM{r}"))       # shape without month / zone / revision words
        vals["BU"] = "=" + alpha_prefix(f"BM{r}")             # letters before the first digit
        vals["BW"] = "=" + digits_only(f"Z{r}")               # digits only
        vals["BX"] = "=" + letters_only(f"BM{r}")             # letters only
        vals["BO"] = "=" + digitmap(f"W{r}")
        vals["BP"] = "=" + alphamap(f"BO{r}")
        vals["CA"] = "=" + alphamap(stripped(f"BO{r}"))
        vals["CB"] = "=" + digits_only(f"W{r}")
        vals["BZ"] = (f"=IFERROR(IF({id_shape_ok(f'Z{r}', f'BN{r}', f'BT{r}', f'BU{r}', f'BW{r}', f'BX{r}', f'S{r}', f'Flights!$F${fr}')},1,0),0)")
        vals["BL"] = f"=IF(ISNA(MATCH(TRIM(S{r}),L_EvidenceType,0)),0,BZ{r})"
        vals["BQ"] = f"=IFERROR(IF({batch_shape_ok(f'W{r}', f'BP{r}', f'CA{r}', f'CB{r}', f'Flights!$F${fr}')},1,0),0)"
        # a plain Pass whose result wording describes a problem is not a Pass (blocks READY; negated forms cancel)
        vals["BG"] = (f"=IFERROR(IF(AND(TRIM(U{r})=\"Pass\",{occurrences('L_ResultBad', f'BC{r}')}>"
                      f"{occurrences('L_NegPhrase', f'BC{r}')}),1,0),0)")

        def bad(x, minlen):
            n = f"{helper[x]}{r}"
            return f"OR(LEN({n})<{minlen},ISNUMBER(MATCH({n},L_Placeholder,0)))"
        if ck.get("due_rule") == "CARRY":
            kl = next((rowmap[x["check_id"]] for x in data["checks"] if x["flight_id"] == ck["flight_id"]
                       and x.get("due_rule") == "CARRY" and x["check_type"] == "Physical uplift" and not x["uplift_stn"]), None)
            cutoff = f"Flights!$AY${fr}+UpliftWindowH/24"
            if kl:
                cutoff = (f"MIN({cutoff},IF(AND(OR(TRIM(U{kl})=\"Pass\",TRIM(U{kl})=\"Pass after CA\"),ISNUMBER(AG{kl})),"
                          f"AG{kl},9E+99))")
            cut_msg = "completed after the items were loaded at KUL / the carrying flight left KUL"
        else:
            cutoff = f"IF(ISNUMBER(Flights!$BA${fr}),MIN(Flights!$BA${fr},Flights!$AA${fr}),Flights!$AA${fr})"
            cut_msg = "completed after loading / departure (cannot establish readiness)"
        if sq and ck["check_type"] == "Preparation":
            cutoff = f"IF({moved},IF(ISNUMBER({carr}),{carr},Flights!$AA${fr}),{cutoff})"
            cut_msg = "completed after loading (at the departure stn, or after the carrying flight left the confirmed stn)"
        msg = (f"IF(AND({U}=\"Pass\",ISNUMBER(AF{r})),IF(AF{r}<>0,\"INVALID {ND} qty variance: use Fail or Pass after CA\",\"\"),\"\")&"
               f"IF(AND({U}=\"Pass after CA\",ISNUMBER(AF{r})),IF(AF{r}<>0,\"INVALID {ND} after the corrective action the actual qty must equal expected (update Actual)\",\"\"),\"\")&"
               f"IF(AND({U}=\"Pass\",TRIM(AB{r})<>\"\",TRIM(AB{r})<>\"Closed\"),\"INVALID {ND} corrective action not closed: use Fail, then Pass after CA\",\"\")&"
               f"IF(AND(LEN(BA{r})>=2,ISNA(MATCH(LEFT(BA{r},255),L_Placeholder,0)),TRIM(AB{r})=\"\"),\"INVALID {ND} CA status missing for the recorded corrective action\",\"\")&"
               f"IF(AND({U}=\"Pass after CA\",OR({bad('AA', 5)},TRIM(AB{r})<>\"Closed\")),\"INVALID {ND} corrective action not recorded/closed\",\"\")")
        rest = (f"{doc_chk}"
                f"IF(AG{r}>AsOfUTC,IF(Settings!$B$4<>\"\",IF(AsOfUTC>Q{r},\"OVERDUE\",\"OPEN\"),\"INVALID {ND} completion time is in the future\"),"
                f"IF(AND(ISNUMBER(BY{r}),AG{r}<BY{r}),IF(H{r}=\"Physical uplift\",\"INVALID {ND} before uplift window (cannot confirm loading)\","
                f"\"INVALID {ND} before valid window\"),"
                f"IF(AG{r}>Flights!$AA${fr},\"INVALID {ND} completed after departure (cannot establish readiness)\","
                f"IF(AND(H{r}=\"Preparation\",AG{r}>={cutoff}),\"INVALID {ND} {cut_msg}\","
                f"IF(AND(H{r}=\"Physical uplift\",AG{r}>Q{r}),\"INVALID {ND} recorded after the loading flight departed\","
                f"IF(AG{r}>Q{r},\"COMPLETE {ND} LATE\",\"COMPLETE\"))))))))))))")
        link_chk, link_close = "", ""
        if ck.get("prep_link"):
            pr = rowmap[ck["prep_link"]]
            link_chk = (f"IF(AH{pr}=\"N/A {ND} JUSTIFIED\",\"INVALID {ND} preparation row {ck['prep_link']} is N/A\","
                        f"IF(AND(AR{r}=1,ISNUMBER(X{pr}),X{r}<>X{pr},TRIM(U{r})<>\"Pass after CA\"),\"INVALID {ND} expected qty differs from preparation check {ck['prep_link']}: record the load change as a corrective action (Pass after CA)\",")
            link_close = "))"
        kul_rows = [rowmap[x["check_id"]] for x in data["checks"]
                     if x["flight_id"] == ck["flight_id"] and x["uplift_stn"] == "KUL" and not x.get("req_carry")]
        carry_follow = bool(ck.get("req_carry")) or (ck.get("due_rule") == "CARRY" and not ck["uplift_stn"])
        kul_all_na = (("AND(" + ",".join(f"LEFT(AH{k},3)=\"N/A\"" for k in kul_rows) + ")")
                      if (kul_rows and carry_follow) else "FALSE")  # only on the follow-on rows (no self-reference)
        written_q = "FALSE" if ck.get("prep_link") else f"N{r}=\"Clarification required\""
        na_link, na_close = "", ""
        if ck.get("prep_link"):
            pr = rowmap[ck["prep_link"]]
            na_link = (f"IF(OR(AH{pr}<>\"N/A {ND} JUSTIFIED\",TRIM(AC{pr})<>TRIM(AC{r})),\"INVALID {ND} preparation line "
                       f"{ck['prep_link']} must also be N/A with the same Outcome (it says carried / is not N/A)\",")
            na_close = ")"
        valid = (
            f"IF(NOT(ISNUMBER(AD{r})),\"INVALID {ND} completion time missing\","
            f"IF(OR({bad('T', 2)},BR{r}<3),\"INVALID {ND} PIC missing or not a named person (a role such as QA / CSM is not a name)\","
            f"IF({bad('V', 2)},\"INVALID {ND} result / assessment missing\","
            f"IF(OR({bad('Z', 3)},{has('L_EvidencePhrase', f'BD{r}')},BI{r}=1),\"INVALID {ND} evidence missing, placeholder or not yet on file\","
            f"IF(ISNA(MATCH(TRIM(S{r}),L_EvidenceType,0)),\"INVALID {ND} choose the evidence type from the list (column S)\","
            f"IF(BL{r}=0,\"INVALID {ND} evidence ID must be the record's own ID only (e.g. SF-2210, IMG_2231, 1045521); a link / file name needs type 'Link / file path'\","
            f"IF(OR({bad('AE', 2)},BS{r}<3),\"INVALID {ND} verifier missing or not a named person (a role such as QA / CSM is not a name)\","
            f"IF(OR(AX{r}=AU{r},CE{r}=1),\"INVALID {ND} verifier must be a different, named person (not the PIC or a role)\","
            f"IF(BH{r}=1,\"INVALID {ND} result not yet available (awaiting / TBC)\","
            f"IF(BG{r}=1,\"INVALID {ND} result describes a problem: record Fail, then Pass after CA (or reword if it is not a problem)\","
            f"IF(AND({written_q},ISNA(MATCH(TRIM(S{r}),{{\"Email\",\"Memo / letter\",\"Document (GLD / menu / ISOP)\",\"System record\",\"Link / file path\"}},0))),\"INVALID {ND} a clarification is resolved only by a written confirmation: evidence type Email, Memo / letter, Document, System record or Link\","
            f"IF(AND(N{r}=\"Clarification required\",TRIM(AC{r})<>INDEX(L_Outcome,1)),\"INVALID {ND} clarification not resolved: choose 'Confirmed – applies / carried as listed' in Outcome (AC) and cite the written confirmation as evidence\","
            f"IF(AND(TRIM(AC{r})<>\"\",TRIM(AC{r})<>INDEX(L_Outcome,1)),\"INVALID {ND} Outcome (AC) says not carried / N/A but the status is Pass\","
            f"IF(OR(LEN(AV{r})<8,LEN(TRIM(BC{r}))-LEN(SUBSTITUTE(TRIM(BC{r}),\" \",\"\"))<1),\"INVALID {ND} result too brief: say what was checked and found (e.g. 'Panel of 3, all 4 dishes to spec')\","
            f"IF(BV{r}=1,\"INVALID {ND} wording says not on board: an on-board confirmation needs the item loaded\","
            f"IF(AND(AQ{r}=1,OR({bad('W', 3)},BJ{r}=1,BQ{r}=0)),\"INVALID {ND} batch ID missing, placeholder or not an ID (e.g. PASB-261008-BC-017)\","
            f"IF(AND(AR{r}=1,OR(NOT(ISNUMBER(X{r})),NOT(ISNUMBER(Y{r})))),\"INVALID {ND} expected/actual qty missing\","
            f"IF(AND(AR{r}=1,OR(X{r}<=0,Y{r}<0)),\"INVALID {ND} expected qty must be above 0 and actual not negative\","
            f"{link_chk}"
            f"IF(BB{r}<>\"\",BB{r},{rest}){link_close}))))))))))))))))))")
        itm = ck["item"].lower()
        for key, dcol, dname in (("gld", "H", "GLD"), ("menu checklist", "M", "menu checklist")):
            if key in itm and not ck["req_doc"]:
                msg += (f"&IF(AND(ISNUMBER(Documents!${dcol}${fr}),Documents!${dcol}${fr}>AD{r}),\"INVALID {ND} {dname} "
                        f"revised after this check (Documents rev date): re-check against the current revision\",\"\")")
        if sq and ck["check_type"] == "Preparation":
            msg += (f"&IF({conf}=\"\",\"INVALID {ND} confirm where this item is loaded (Settings, station confirmations {sq['qid']})\","
                    f"IF(AND({moved},NOT(ISNUMBER({carr}))),\"INVALID {ND} enter the carrying flight's departure from the confirmed stn (Settings {sq['qid']})\","
                    f"IF(AND({moved},OR({arr}>Flights!$AA${fr},{carr}<Flights!$AA${fr}-3)),\"INVALID {ND} carrying flight ({sq['qid']}) must arrive before STD and leave within 3 days of it\",\"\")))")
        if "Qty per tail: A350 280 pcs" in (ck["note"] or ""):
            tq = (f"IF(UPPER(SUBSTITUTE(SUBSTITUTE(TRIM(Flights!$N${fr}),\"-\",\"\"),\" \",\"\"))=\"9MMAH\",280,260)")
            msg += (f"&IF(AND(ISNUMBER(X{r}),X{r}<>{tq}),\"INVALID {ND} expected qty must be the reference quantity for this "
                    f"tail: 280 for 9M-MAH, 260 for an A359 (check Tail/Reg on Flights)\",\"\")")
        vals["BB"] = "=" + msg
        if ck.get("prep_link"):
            pq = rowmap[ck["prep_link"]]
            clar_q = (f"AND(N{r}=\"Clarification required\",TRIM(AC{r})=\"\",TRIM(AC{pq})=\"\","
                      f"LEFT(AH{pq},8)<>\"COMPLETE\",AH{pq}<>\"N/A {ND} JUSTIFIED\")")
        else:
            clar_q = f"AND(N{r}=\"Clarification required\",TRIM(AC{r})=\"\")"
        state = (
            f"=IF(N{r}=\"N/A {ND} rule\",IF(OR({U}=\"\",{U}=\"N/A\",{U}=\"Not started\"),\"N/A {ND} RULE\","
            f"\"INVALID {ND} this check does not apply to this aircraft: clear the status (or correct the fleet on Flights)\"),"
            f"IF({U}=\"N/A\",IF(AT{r}=0,\"INVALID {ND} N/A not permitted for this check\","
            f"IF(AND(TRIM(AB{r})<>\"\",TRIM(AB{r})<>\"Closed\"),\"INVALID {ND} corrective action still open\","
            f"IF(OR(AND(AT{r}<>2,TRIM(AC{r})<>INDEX(L_Outcome,2),TRIM(AC{r})<>INDEX(L_Outcome,3)),"
            f"AND(AT{r}=2,TRIM(AC{r})<>INDEX(L_Outcome,4))),"
            f"\"INVALID {ND} N/A needs its reason chosen in Outcome (AC): \"&IF(AT{r}=2,\"'Refreshment service – no printed menu card'\","
            f"\"'Confirmed – not carried on this sector' or '… not applicable to this aircraft / tail'\"),"
            f"IF(AND(AT{r}<>2,OR(BL{r}=0,ISNA(MATCH(TRIM(S{r}),{{\"Email\",\"Memo / letter\",\"Document (GLD / menu / ISOP)\",\"System record\",\"Link / file path\"}},0)))),\"INVALID {ND} N/A on a clarification needs the written confirmation as evidence (Email, Memo / letter, Document, System record or Link, with its ID)\","
            f"IF(AND(AT{r}=3,NOT({kul_all_na})),\"INVALID {ND} N/A here only once every KUL-sourced item of this flight is confirmed not carried (N/A)\","
            f"IF(NOT(ISNUMBER(AD{r})),\"INVALID {ND} completion time missing (when the N/A was confirmed)\","
            f"IF(AG{r}>AsOfUTC,IF(Settings!$B$4<>\"\",IF(AsOfUTC>Q{r},\"OVERDUE\",\"OPEN\"),\"INVALID {ND} completion time is in the future\"),"
            f"IF(AG{r}>Flights!$AA${fr},\"INVALID {ND} N/A confirmed after departure (cannot establish readiness)\","
            f"{na_link}"
            f"IF(AND(NOT({bad('T', 2)}),BR{r}>=3,NOT({bad('AE', 2)}),BS{r}>=3,"
            f"TRIM(AE{r})<>TRIM(T{r}),AX{r}<>AU{r},CE{r}=0),\"N/A {ND} JUSTIFIED\","
            f"\"INVALID {ND} N/A needs a named PIC and a different named verifier\")))))))){na_close}),"
            f"IF(OR({U}=\"Pass\",{U}=\"Pass after CA\"),{valid},"
            f"IF({U}=\"Fail\",\"FAIL {ND} DISCREPANCY\","
            f"IF(AND({U}<>\"\",{U}<>\"Not started\",{U}<>\"In progress\"),\"INVALID {ND} unrecognised status (use the list)\","
            f"IF(AsOfUTC>Q{r},\"OVERDUE\",IF({clar_q},\"OPEN {ND} CLARIFICATION\",\"OPEN\")))))))")
        refresh_menu = ck["category"] == "Menu" and "Refreshment service" in ck["note"]
        na_ok = (2 if refresh_menu else 3 if carry_follow
                 else int(ck["applic"] == "Clarification required"))
        vals["AT"] = na_ok
        vals["AH"] = f"=IFERROR({state[1:]},\"INVALID {ND} error value in an input cell\")"
        vals["AI"] = f"=IF(OR(AH{r}=\"N/A {ND} RULE\",AH{r}=\"N/A {ND} JUSTIFIED\"),0,1)"
        vals["AJ"] = f"=IF(LEFT(AH{r},8)=\"COMPLETE\",1,0)"
        vals["AK"] = f"=IF(AND(AI{r}=1,AJ{r}=0,AsOfUTC>Q{r}),1,0)"
        vals["AL"] = (f"=IFERROR(IF(AND(AI{r}=1,OR(TRIM(U{r})=\"Fail\",AND(TRIM(AB{r})<>\"\",TRIM(AB{r})<>\"Closed\"),AND(ISNUMBER(AF{r}),AF{r}<>0,TRIM(U{r})<>\"Pass after CA\"))),1,0),1)")
        vals["AM"] = f"=IF(LEFT(AH{r},7)=\"INVALID\",1,0)"
        vals["AN"] = f"=IF(AND({clar_q},AI{r}=1,AJ{r}=0),1,0)"
        vals["AO"] = f"=IF(AK{r}=1,SUM(AK$5:AK{r}),\"\")"
        vals["AP"] = f"=IF(AL{r}=1,SUM(AL$5:AL{r}),\"\")"
        for col, v in vals.items():
            ws[f"{col}{r}"] = v
        height = 1
        for j in range(1, len(CK_HEAD) + 1):
            c = ws.cell(r, j)
            c.font = f(9)
            c.border = BORDER
            c.alignment = WRAP
            if 20 <= j <= 31 or j == 19:
                c.fill = F_INPUT
            elif j >= 32:
                c.fill = F_CALC
            elif j >= 11 and j <= 15:
                c.fill = F_REF
        if exp_qty is not None:
            ws[f"X{r}"].fill = F_REF  # reference quantity: locked (not yellow) so it cannot be lowered
        if isinstance(exp_qty, str):
            ws[f"X{r}"].font = f(9, color="0000FF")
            ws[f"X{r}"].comment = Comment("Follows the tail on Flights N: 9M-MAH = 280 pcs, other A350 (A359) = 260 pcs "
                                          "(SEAT LINEN!B76:C77). Blank until the tail is entered.", "MAGCS")
        elif exp_qty is not None:
            ws[f"X{r}"].font = f(9, color="0000FF")
            ws[f"X{r}"].comment = Comment(f"Pre-filled from reference: {ck['src']}", "MAGCS")
        for col in ("D",):
            ws[f"{col}{r}"].number_format = "dd-mmm-yy"
        for col in ("Q", "R", "BY", "AD", "AG"):
            ws[f"{col}{r}"].number_format = DT
        ws[f"AH{r}"].font = f(9, True)
        if not ck["req_batch"]:
            ws[f"W{r}"].fill = F_CALC
        for col, w in (("J", CK_W[9]), ("K", CK_W[10]), ("L", CK_W[11]), ("O", CK_W[14])):
            height = max(height, est_lines(ws[f"{col}{r}"].value, w))
        ws.row_dimensions[r].height = max(40, 12.5 * height + 4)
    # validations
    dv = DataValidation(type="list", formula1="=L_Status", allow_blank=True,
                        prompt="Pass needs completion time, evidence and verifier. N/A needs its reason in Outcome (AC) and a verifier.",
                        promptTitle="Status", showInputMessage=True, showErrorMessage=True)
    dv.add(f"U5:U{last}")
    ws.add_data_validation(dv)
    dvt = DataValidation(type="list", formula1="=L_EvidenceType", allow_blank=True, showErrorMessage=True,
                         showInputMessage=True, promptTitle="Evidence type",
                         prompt="Choose the kind of record, then enter ONLY its ID in column Z (e.g. SF-2210, IMG_2231, "
                                "CO/26/331, 1045521), or a link / file path.", errorTitle="Evidence type",
                         error="Choose an evidence type from the list.")
    dvt.add(f"S5:S{last}")
    ws.add_data_validation(dvt)
    dvo = DataValidation(type="list", formula1="=L_Outcome", allow_blank=True, showErrorMessage=True,
                         showInputMessage=True, promptTitle="Outcome",
                         prompt="Clarification lines: choose how it was resolved (cite the written confirmation as evidence). "
                                "N/A: choose the reason.", errorTitle="Outcome", error="Choose an outcome from the list.")
    dvo.add(f"AC5:AC{last}")
    ws.add_data_validation(dvo)
    dv2 = DataValidation(type="list", formula1="=L_CA", allow_blank=True, showErrorMessage=True,
                         errorTitle="CA status", error="Choose Open or Closed from the list.")
    dv2.add(f"AB5:AB{last}")
    ws.add_data_validation(dv2)
    dv3 = DataValidation(type="decimal", operator="between", formula1="46204", formula2="46419", allow_blank=True,
                         showErrorMessage=True, showInputMessage=True, promptTitle="Completion time",
                         prompt="Local time at the check station (column P), e.g. 01/10/2026 14:30.",
                         errorTitle="Date-time required", error="Enter a date-time between Jul and Dec 2026.")
    dv3.add(f"AD5:AD{last}")
    ws.add_data_validation(dv3)
    dv4 = DataValidation(type="decimal", operator="greaterThanOrEqual", formula1="0", allow_blank=True,
                         showErrorMessage=True, error="Quantity must be a number >= 0.")
    dv4.add(f"X5:Y{last}")
    ws.add_data_validation(dv4)
    # conditional formatting
    st = f"AH5:AH{last}"
    green = PatternFill("solid", fgColor="C6EFCE")
    red = PatternFill("solid", fgColor="FFC7CE")
    amber = PatternFill("solid", fgColor="FFEB9C")
    grey = PatternFill("solid", fgColor="D9D9D9")
    ws.conditional_formatting.add(st, FormulaRule(formula=['ISNUMBER(SEARCH("CHECK WORDING",AH5))'], fill=amber, stopIfTrue=True))
    ws.conditional_formatting.add(st, FormulaRule(formula=['LEFT(AH5,8)="COMPLETE"'], fill=green))
    ws.conditional_formatting.add(st, FormulaRule(formula=['OR(LEFT(AH5,7)="INVALID",AH5="OVERDUE",LEFT(AH5,4)="FAIL")'], fill=red))
    ws.conditional_formatting.add(st, FormulaRule(formula=['LEFT(AH5,4)="OPEN"'], fill=amber))
    ws.conditional_formatting.add(st, FormulaRule(formula=['LEFT(AH5,3)="N/A"'], fill=grey))
    ws.conditional_formatting.add(f"A5:S{last}", FormulaRule(formula=['$AH5="N/A – RULE"'], font=Font(name=FONT, color="808080")))
    ws.conditional_formatting.add(f"N5:N{last}", CellIsRule(operator="equal", formula=['"Clarification required"'],
                                  font=Font(name=FONT, color="C00000", bold=True)))
    ws.freeze_panes = "K5"
    ws.auto_filter.ref = f"A4:AS{last}"
    ws.print_area = f"A1:AH{last}"
    ws.print_title_cols = "A:C"
    ws.sheet_view.zoomScale = 85
    for col in ("AQ", "AR", "AS", "AT", "AU", "AV", "AW", "AX", "AY", "AZ", "BA", "BB", "BC", "BD", "BE", "BF", "BG", "BH", "BI", "BJ", "BK", "BL", "BM", "BN", "BO", "BP", "BQ", "BR", "BS", "BT", "BU", "BV", "BW", "BX", "BY", "BZ", "CA", "CB", "CC", "CD", "CE"):
        ws.column_dimensions[col].hidden = True
    fit_pages(ws, "A", "AH", title_cols_w=29)
    ws.print_title_rows = "4:4"
    return n


# ------------------------------------------------------------------ Documents
def build_documents(wb, data):
    ws = wb.create_sheet("Documents")
    title(ws, "Documents",
          "PIC reference documents: galley loading diagrams (GLD) and menu checklists. No GLD or menu checklist was supplied with the brief: every row starts OUTSTANDING. A row is ON FILE only with a real doc no, revision, a revision date between 2020 and the flight date, and an attachment location. Enter doc no, revision, "
          "revision date and attachment location/link (or embed on the flight's P-sheet) to clear it.",
          "C2:I2", 99)
    labels = ["Flight ID", "Flight No", "Date", "Sector", "Fleet", "GLD doc no", "GLD revision", "GLD rev date (issued or received, whichever later)",
              "GLD attachment (link / location)", "GLD status", "Menu checklist doc no", "Menu checklist revision",
              "Menu checklist rev date (issued or received, whichever later)", "Menu checklist attachment (link / location)", "Menu checklist status",
              "Outstanding", "Notes"]
    widths = [7, 9, 10, 9, 9, 16, 10, 11, 34, 14, 16, 10, 11, 34, 14, 10, 40]
    header(ws, 4, labels, widths, height=42)
    for i, fl in enumerate(data["flights"]):
        r = 5 + i
        fr = r
        ws[f"A{r}"] = fl["id"]
        ws[f"B{r}"] = f"=Flights!F{fr}"
        ws[f"C{r}"] = f"=Flights!D{fr}"
        ws[f"C{r}"].number_format = "dd-mmm-yy"
        ws[f"D{r}"] = f"=Flights!G{fr}&\"-\"&Flights!H{fr}"
        ws[f"E{r}"] = f"=Flights!L{fr}"
        def docok(no, rev, dt, att):
            ok = lambda x, n: f"AND(LEN({norm(x + str(r))})>={n},NOT(ISNUMBER(MATCH(LEFT({norm(x + str(r))},255),L_Placeholder,0))))"
            docno = (f"AND(SUMPRODUCT(--ISNUMBER(FIND({{\"0\",\"1\",\"2\",\"3\",\"4\",\"5\",\"6\",\"7\",\"8\",\"9\"}},{no}{r})))>0,"
                     f"ISNA(MATCH(TRIM({no}{r}),{{\"12345\",\"123456\",\"00000\",\"000000\",\"11111\",\"99999\"}},0)),"
                     f"LOWER(TRIM({no}{r}))<>LOWER(Flights!$F${fr}),NOT(AND(LEFT(LOWER(TRIM({no}{r})),2)=\"mh\",ISNUMBER(--MID(TRIM({no}{r}),3,6)))))")
            return (f"=IFERROR(IF(AND({ok(no, 3)},{docno},{ok(rev, 1)},ISNUMBER({dt}{r}),{dt}{r}>=DATE(2020,1,1),"
                    f"{dt}{r}<=Flights!$I${fr},{ok(att, 5)},OR({ref_ok(att + str(r), ('R' if att == 'I' else 'T') + str(r), ('S' if att == 'I' else 'U') + str(r))},SUMPRODUCT(--ISNUMBER(SEARCH({{\"binder\",\"cabinet\",\"shelf\",\"filed in\",\"file room\",\"lever arch\"}},{att}{r})))>0),NOT({pending(('X' if att == 'I' else 'Y') + str(r), 'L_PendStrict')})),\"ON FILE\",\"OUTSTANDING\"),\"OUTSTANDING\")")
        ws[f"X{r}"] = "=" + spaced(f"I{r}")
        ws[f"Y{r}"] = "=" + spaced(f"N{r}")
        ws[f"V{r}"] = "=" + masked(f"X{r}", f"Flights!$F${fr}", 1)
        ws[f"W{r}"] = "=" + masked(f"Y{r}", f"Flights!$F${fr}", 1)
        ws[f"Z{r}"] = "=" + masked(f"V{r}", f"Flights!$F${fr}", 2)
        ws[f"AA{r}"] = "=" + masked(f"W{r}", f"Flights!$F${fr}", 2)
        ws[f"AB{r}"] = "=" + masked(f"Z{r}", f"Flights!$F${fr}", 3)
        ws[f"AC{r}"] = "=" + masked(f"AA{r}", f"Flights!$F${fr}", 3)
        ws[f"R{r}"] = "=" + digitmap(f"AB{r}")
        ws[f"S{r}"] = "=" + alphamap(f"R{r}")
        ws[f"T{r}"] = "=" + digitmap(f"AC{r}")
        ws[f"U{r}"] = "=" + alphamap(f"T{r}")
        ws[f"J{r}"] = docok("F", "G", "H", "I")
        ws[f"O{r}"] = docok("K", "L", "M", "N")
        ws[f"P{r}"] = f"=(J{r}=\"OUTSTANDING\")+(O{r}=\"OUTSTANDING\")"
        note = "Not supplied with brief – obtain from caterer / MAGCS."
        if fl["round_trip"]:
            note += " Return catering loaded at KUL: GLD must show return-leg stowage."
        ws[f"Q{r}"] = note
        for j in range(1, 18):
            c = ws.cell(r, j)
            c.font = f(9)
            c.border = BORDER
            c.alignment = WRAP
            if j in (6, 7, 8, 9, 11, 12, 13, 14):
                c.fill = F_INPUT
            elif j in (10, 15, 16):
                c.fill = F_CALC
        ws[f"H{r}"].number_format = "dd-mmm-yy"
        ws[f"M{r}"].number_format = "dd-mmm-yy"
        for col in ("A", "C", "D", "E", "J", "O", "P"):
            ws[f"{col}{r}"].alignment = Alignment(horizontal="center", vertical="top", wrap_text=True)
        ws.row_dimensions[r].height = 30
    for rng in ("J5:J26", "O5:O26"):
        ws.conditional_formatting.add(rng, CellIsRule(operator="equal", formula=['"OUTSTANDING"'],
                                      fill=PatternFill("solid", fgColor="FFC7CE"), font=Font(name=FONT, color="9C0006", bold=True)))
        ws.conditional_formatting.add(rng, CellIsRule(operator="equal", formula=['"ON FILE"'],
                                      fill=PatternFill("solid", fgColor="C6EFCE"), font=Font(name=FONT, color="006100", bold=True)))
    dvd = DataValidation(type="date", operator="between", formula1="43831", formula2="46387", allow_blank=True,
                         showErrorMessage=True, error="Enter the revision date (2020–2026).")
    dvd.add("H5:H26")
    dvd.add("M5:M26")
    ws.add_data_validation(dvd)
    ws.freeze_panes = "C5"
    ws.print_title_cols = "A:B"
    ws.print_title_rows = "4:4"
    for col in ("R", "S", "T", "U", "V", "W", "X", "Y", "Z", "AA", "AB", "AC"):
        ws.column_dimensions[col].hidden = True
    ws.print_area = "A1:Q26"
    fit_pages(ws, "A", "Q", title_cols_w=16)


def build_isop(wb):
    ws = wb.create_sheet("ISOP Register")
    title(ws, "ISOP revision register",
          "Record every applicable In-flight Service Operating Procedure revision communicated to caterers. "
          "The T-7D ISOP checks on each flight should quote the revision numbers listed here.", "A2:I2", 140)
    header(ws, 4, ["#", "ISOP doc / section", "Revision", "Effective date", "Title / change summary", "Applies to stations",
                   "Communicated on", "Communicated by", "Caterer acknowledgement ref"],
           [5, 16, 9, 11, 30, 14, 14, 16, 20], height=32)
    for k in range(40):
        r = 5 + k
        ws.cell(r, 1, k + 1).font = f(9)
        for j in range(1, 10):
            c = ws.cell(r, j)
            c.border = BORDER
            c.font = f(9)
            c.alignment = WRAP
            if j > 1:
                c.fill = F_INPUT
        ws.cell(r, 4).number_format = "dd-mmm-yy"
        ws.cell(r, 7).number_format = "dd-mmm-yy"
        ws.row_dimensions[r].height = 24
    ws.freeze_panes = "B5"
    ws.print_title_rows = "4:4"
    fit_pages(ws, "A", "I", landscape=True)


def fit_pages(ws, first, last, landscape=True, title_cols_w=0):
    """Pages-wide so that print scale stays >= ~85% (A4: ~148 units landscape, ~103 portrait)."""
    from openpyxl.utils import column_index_from_string as ci
    total = sum((ws.column_dimensions[get_column_letter(i)].width or 8.43)
                for i in range(ci(first), ci(last) + 1)
                if not ws.column_dimensions[get_column_letter(i)].hidden)
    per = (148 if landscape else 103) * 0.92 - title_cols_w
    pages = max(1, -(-int(total - title_cols_w) // int(per)))
    ws.page_setup.orientation = "landscape" if landscape else "portrait"
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.fitToWidth = pages
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_margins.left = ws.page_margins.right = 0.3
    ws.oddFooter.center.text = "&A  |  Page &P of &N"
    ws.oddFooter.center.size = 8


# ------------------------------------------------------------------ Requirements
def build_requirements(wb, data):
    ws = wb.create_sheet("Requirements")
    title(ws, "Requirements",
          "Uplift requirements by flight, derived from STD UPLIFT INFORMATION.xlsx. Applied by sector, aircraft, assessed cabin class and uplift station. 'Not stated in reference' = the reference "
          "gives no uplift station for that item. Clarification rows block READY until resolved (Pass or justified N/A).",
          "C2:I2", 108)
    labels = ["Flight ID", "Flight No", "Date", "Sector", "Class", "Fleet", "Category", "Item", "Expected",
              "Qty", "Uplift stn", "Applicability", "Note / open question", "Source cell"]
    widths = [8, 11, 10, 9, 6, 9, 14, 30, 30, 12, 14, 14, 50, 38]
    header(ws, 4, labels, widths)
    byid = {x["id"]: x for x in data["flights"]}
    r = 5
    for q in data["requirements"]:
        fl = byid[q["flight_id"]]
        vals = [fl["id"], fl["flt"], dt(fl["date"] + " 00:00"), f"{fl['dep']}-{fl['arr']}", fl["cls"], fl["fleet"],
                q["category"], q["item"], q["expected"], q["qty"], q["uplift_stn"], q["applic"], q["note"], q["src"]]
        h = 1
        for j, v in enumerate(vals, 1):
            c = ws.cell(r, j, v)
            c.font = f(9, bold=(j == 12 and v != "Required"), color="C00000" if (j == 12 and v != "Required") else "000000")
            c.border = BORDER
            c.alignment = WRAP
            h = max(h, est_lines(v, widths[j - 1]))
        ws.cell(r, 3).number_format = "dd-mmm-yy"
        ws.row_dimensions[r].height = max(15, 12.5 * h + 3)
        r += 1
    r += 1
    ws.cell(r, 1, "Items the reference marks NOT provided (X / blank) for each flight – listed for traceability, no check line created").font = f(11, True, NAVY)
    r += 1
    header(ws, r, ["Flight ID", "Flight No", "Date", "Sector", "Class", "Fleet", "Category", "Item", "Source cell"], height=30)
    r += 1
    for q in data["not_provided"]:
        fl = byid[q["flight_id"]]
        vals = [fl["id"], fl["flt"], dt(fl["date"] + " 00:00"), f"{fl['dep']}-{fl['arr']}", fl["cls"], fl["fleet"],
                q["category"], q["item"], q["src"]]
        for j, v in enumerate(vals, 1):
            c = ws.cell(r, j, v)
            c.font = f(9, color="595959")
            c.border = BORDER
            c.alignment = WRAP
        ws.cell(r, 3).number_format = "dd-mmm-yy"
        ws.row_dimensions[r].height = max(15, 12.5 * est_lines(q["src"], 30) + 3)
        r += 1
    ws.freeze_panes = "C5"
    ws.auto_filter.ref = f"A4:N{4 + len(data['requirements'])}"
    ws.print_title_rows = "4:4"
    ws.print_title_cols = "A:B"
    fit_pages(ws, "A", "N", title_cols_w=19)


# ------------------------------------------------------------------ Instructions
INSTR = [
    ("h", "How to use this workbook"),
    ("p", "Purpose: monitor MAGCS inflight-catering readiness for all 22 legs of the Skytrax Agenda 2026 audit itineraries, "
          "with three checkpoints per flight calculated from scheduled departure (STD) in the correct station time zone: T-7D, T-24H and T-12H. "
          "T-12H is recorded in two parts: the preparation check at the caterer (T-12H PREP) and the physical on-board uplift confirmation "
          "before departure (UPLIFT), so dashboards show four columns."),
    ("h2", "Sheets"),
    ("b", "Dashboard – checkpoint completion, overdue actions, outstanding discrepancies and readiness per flight."),
    ("b", "Flights – the 22 legs as scheduled. Enter Tail/Reg, Caterer and Flight PIC (yellow). Due times: T-7D, T-24H, T-12H prep and uplift (= STD)."),
    ("b", "Checks – one row per check per flight. THE ONLY PLACE TO RECORD RESULTS (yellow columns S–AE)."),
    ("b", "Documents – galley loading diagram and menu checklist per flight (doc no, revision, date, attachment). ISOP Register – ISOP revisions communicated to caterers."),
    ("b", "Requirements – uplift requirements derived from STD UPLIFT INFORMATION.xlsx per flight, with source cell for each."),
    ("b", "Settings – as-of time, PC clock offset, validity windows, station time-zone table."),
    ("b", "P01–P22 – printable pack per flight (formula views of Checks, A4 landscape). Page 1 is a cover sheet with the flight's due times, readiness, the attachments register (A1 galley loading diagram, A2 menu checklist – staple behind, flagged OUTSTANDING until on file) and sign-off; the checklist starts on page 2. Print one flight's sheet on its own so page numbers run per flight."),
    ("h2", "Recording a check (Checks sheet)"),
    ("n", "1. Filter column B (Flight ID) or G (Checkpoint). Enter PIC (T)."),
    ("n", "2. Record Result (V), Evidence type (S) and Evidence ID (Z), Completion time in LOCAL time of the check station shown in column P (AD) and Verifier (AE)."),
    ("n", "3. T-24H sensory: Batch ID (W) is mandatory. Status Pass = batch accepted; Fail = rejected (enter corrective action AA and CA status AB)."),
    ("n", "4. Quantity lines: enter Expected (X, from menu checklist / GLD) and Actual (Y). A variance cannot be 'Pass' – use Fail, then 'Pass after CA' once the corrective action is recorded and Closed."),
    ("n", "5. Set Status (U) last. The Record state (AH) tells you whether the entry is accepted."),
    ("n", "6. N/A only with a written justification (AC), PIC (T) and verifier (AE), and only on rows where N/A is permitted. Rule-based N/A (digital IFE menu on non-A339) is automatic and excluded from completion rates."),
    ("h2", "Rules built into the formulas"),
    ("b", "Blank or 'Not started'/'In progress' never counts as complete. 'Pass' without completion time, PIC, evidence or verifier (blank or spaces) shows INVALID and does not count."),
    ("b", "A check completed after the flight's departure never counts (INVALID – completed after departure). A check completed after its due time but before departure counts as COMPLETE – LATE."),
    ("b", "N/A is only accepted where it can legitimately apply: printed menu cards on refreshment-only flights, clarification items confirmed in writing, and KUL-loading lines once every KUL-sourced item is confirmed not carried. Mandatory checks (sensory tests, catering officer, ISOP, GLD, menu checklist, meal/equipment preparation and physical uplift) cannot be N/A'd."),
    ("b", "Completion % = complete ÷ in-scope checks; justified and rule-based N/A are removed from both, so they neither raise nor lower the rate."),
    ("b", "Preparation checks (T-12H PREP, at the caterer) never confirm loading. Physical uplift rows (UPLIFT) are only accepted when the completion time is inside the uplift window before STD (Settings B7) – an earlier entry shows 'INVALID – before uplift window'."),
    ("b", "Completion times in the future, before the valid window or after departure (uplift) are rejected."),
    ("b", "GLD / menu-checklist checks cannot be passed until the Documents row is ON FILE (doc no, revision, date and attachment all present). The revision date is the date the revision was issued or received at the station, whichever is later: every line checked 'vs GLD' or 'vs menu checklist' before that date is INVALID and must be re-checked against the current revision."),
    ("b", "Readiness order: NOT READY – OVERDUE, – DISCREPANCY, – INVALID ENTRY, – DOCUMENTS (GLD / menu checklist not on file), – CLARIFICATION (reference question open), then PREP DONE – AWAITING UPLIFT, IN PROGRESS or NOT STARTED."),
    ("b", "READY only when every in-scope check (including physical uplift) is complete, with zero open discrepancies, zero invalid entries and both documents on file. 'Clarification required' rows (reference ambiguous) also block READY until confirmed (Pass) or justified N/A."),
    ("b", "Overdue = not complete and the effective as-of time is past the due time. Settings B4 is an optional override (UTC); when blank the live clock is used. The effective as-of time is shown in Settings B6."),
    ("b", "Sheets are protected without a password so formulas cannot be overtyped by accident; yellow input cells stay editable and filtering, row sizing and inserting pictures still work. Do not sort the Checks sheet – the P-sheets read fixed rows; use the filters instead. Review > Unprotect Sheet if a structural change is needed."),
    ("b", "Evidence is recorded in two fields: choose the Evidence type in column S, then enter ONLY the record's own ID in column Z (describe the record in Result, not here). An ID is one token with letters joined to at least two digits (SF-2210, IMG_2231, LS/OCT/0118, DN88213, PCS-0727) or a number of 5+ digits (seal 1045521). Links, network paths and file names (…\\SF-2210.pdf, https://…) only with type 'Link / file path'. Rejected: flight numbers, dates, times, revisions ('Rev3'), aircraft types, seat or tray numbers, several IDs in one cell (give the main one) and dummy values (00000, 12345, XX-0000). Batch IDs likewise (PASB-261008-BC-017 or 261008-017). Document numbers on the Documents sheet need digits and must not be the flight number. PIC and verifier must be named people – a role or title alone ('QA', 'PASB supervisor', 'Station Manager KUL') is rejected; Settings V lists the role words."),
    ("b", "Write results with detail: one-word entries such as 'Good', 'Confirmed', 'Checked OK' or 'Self' are rejected. A plain Pass whose result describes a problem ('3 trays missing', 'menu cards not delivered', '2 meals short') is INVALID and blocks READY: record Fail, then 'Pass after CA' with the corrective action. Negated or zero counts ('no defects found', '0 discrepancies', 'nothing missing') are fine; if a genuine pass is flagged (e.g. a dish name), reword it. Settings M–N hold the word lists."),
    ("b", "Placeholder text (e.g. '-', '?', 'TBC', 'n/a', 'pending' – list on Settings K) never counts as evidence, PIC, verifier, result or batch ID. Results need at least two words saying what was checked and found – stock phrases such as 'All good', 'As per menu' or 'Conforms' are rejected. OUTCOME (column AC, dropdown – Settings X): a 'Clarification required' line is resolved only by choosing 'Confirmed – applies / carried as listed' (then Pass) or 'Confirmed – not carried on this sector' / '… not applicable to this aircraft / tail' (then N/A), and citing the written confirmation (e.g. the MAGCS email) as evidence type + ID. Printed menu cards on refreshment-only flights may be N/A with 'Refreshment service – no printed menu card'. N/A also needs a named PIC and a different named verifier and is not permitted anywhere else. A status entered on a check that does not apply to the aircraft (grey 'N/A – RULE' row) is flagged INVALID – clear it."),
    ("b", "Preparation checks must be completed before the catering is loaded: before the first on-board confirmation for the flight, or for KUL-loaded items before the carrying flight leaves KUL. On-board checks at an outstation are only valid once the carrying flight could have arrived."),
    ("b", "Quantity lines: after a corrective action, update Actual to the corrected quantity; 'Pass after CA' requires Actual = Expected."),
    ("b", "PIC and verifier must be different people. Every Pass needs a result/assessment. T-7D evidence older than the Settings B10 window is rejected as stale."),
    ("b", "Row heights do not grow automatically for long typed entries: after entering long text, select the rows and use Home > Format > AutoFit Row Height."),
    ("h2", "Time zones & special cases"),
    ("b", "Due times are computed from STD converted to UTC, then shown in local time of the check station. LHR changes from BST to GMT on 25-Oct-2026; ADL is on ACDT (UTC+10:30) from 4-Oct-2026."),
    ("b", "MH1149 PEN-KUL and MH1437 LGK-KUL: the reference lists KUL as the only uplift station, so all their catering is loaded at KUL on a carrying flight (agenda candidates MH1140 / MH1450 shown in Flights W). Enter the actual carrying flight's KUL departure (UTC) in Flights AX: T-24H, T-12H prep and the KUL loading-confirmation line are capped at that departure minus the uplift window. Until entered they assume the agenda candidate's KUL departure."),
    ("b", "Outstation departures that carry KUL-sourced items (e.g. MH0003 LHR-KUL: pajamas, slippers, signature drinks) have a clarification check to identify the inbound KUL flight. Enter its KUL departure (UTC) on Flights column AX: the preparation due and the KUL loading-confirmation line (UPLIFT, station KUL) follow it. Until entered, they assume the latest possible KUL departure."),
    ("b", "Items whose uplift station differs from the departure station (e.g. signature drinks and slippers for LHR-KUL and HKG-KUL are uplifted at KUL) show the item uplift station in Checks column M. Their preparation due time is capped at the KUL loading deadline on Flights AY (carrying flight's KUL departure in AX minus the uplift window)."),
    ("h2", "Outstanding inputs at issue"),
    ("b", "Galley loading diagrams and menu checklists were not supplied – all 44 are flagged OUTSTANDING (Documents sheet, Dashboard and each P-sheet cover). Record doc no, revision, revision date and file location on Documents; in Excel you may also insert the image in the area after each P-sheet's checklist."),
    ("b", "A350 tail (A359 vs 9M-MAH) decides sales-cart location and EY blanket quantity – enter Tail/Reg on Flights (9M-MAB … 9M-MAH). An A350 flight cannot be READY until its tail is entered (counted as an open clarification)."),
    ("b", "Open uplift-station questions (A350 toiletry kits on KUL-LHR legs, sales carts on return legs – 7 in all) are answered on Settings, 'Station confirmations' (row 26 onward): choose the CONFIRMED uplift stn, and if it is not the departure station enter the carrying flight's departure from that station (UTC). The preparation line then uses that station's time zone, due time and loading cap (it must be done before the carrying flight leaves), and the on-board line cannot be confirmed before the carrier arrives. Until answered, the preparation line cannot be passed."),
    ("b", "Caterer per station, ISOP revision numbers and expected meal/equipment quantities come from the caterer/MAGCS documents – not invented here."),
    ("h2", "Illustrative entry (example only – not recorded anywhere in this workbook)"),
    ("p", "Check F02-T24-01 · PIC: A. Rahman · Status: Pass · Result: Panel score 4.5/5, texture & temperature OK · Batch ID: PASB-261008-BC-017 · "
          "Evidence: Sensory panel sheet · ID SENS-261008-017 · Completion time: 08-Oct-26 20:10 (KUL) · Verifier: N. Ismail"),
    ("h2", "Colour key"),
    ("k1", "Yellow cell = input – enter data here"),
    ("k2", "Grey cell = formula – do not overwrite"),
    ("k3", "Light blue cell = requirement text derived from STD UPLIFT INFORMATION.xlsx"),
    ("k4", "Blue text = value transcribed from a reference (schedule, time-zone table, reference quantity)"),
]


def build_instructions(wb, data):
    ws = wb.create_sheet("Instructions")
    ws.column_dimensions["A"].width = 3
    ws.column_dimensions["B"].width = 88
    ws["B1"] = "MAGCS Inflight Catering – Skytrax 2026 Readiness Checklist & Monitor"
    ws["B1"].font = f(16, True, NAVY)
    ws.row_dimensions[1].height = 26
    ws["B2"] = (f"Built from: Skytrax Agenda 2026 (22 legs, 3 itineraries) and STD UPLIFT INFORMATION.xlsx "
                f"(aircraft data: {data['ref_revision']}).")
    ws["B2"].font = f(9, italic=True, color="595959")
    r = 4
    for kind, text in INSTR:
        c = ws.cell(r, 2, ("•  " + text) if kind == "b" else text)
        if kind == "h":
            c.font = f(13, True, NAVY)
        elif kind == "h2":
            c.font = f(11, True, NAVY)
            ws.row_dimensions[r].height = 22
            c.alignment = Alignment(vertical="bottom")
        else:
            c.font = f(10)
            c.alignment = WRAP
            ws.row_dimensions[r].height = max(15, 13.5 * est_lines(c.value, 74) + 3)
        if kind == "k1":
            c.fill = F_INPUT
        elif kind == "k2":
            c.fill = F_CALC
        elif kind == "k3":
            c.fill = F_REF
        elif kind == "k4":
            c.font = f(10, color="0000FF")
        r += 1
    ws.sheet_view.showGridLines = False
    fit_pages(ws, "A", "B", landscape=False)


def build(wb, data):
    build_instructions(wb, data)
    build_settings(wb, data)
    n = len(data["checks"])
    build_flights(wb, data, n)
    build_checks(wb, data)
    build_documents(wb, data)
    build_requirements(wb, data)
    build_isop(wb)

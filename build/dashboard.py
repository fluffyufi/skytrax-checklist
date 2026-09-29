"""Dashboard sheet for the MAGCS Skytrax 2026 catering readiness workbook.

Every figure on this sheet is a formula that reads Flights / Checks /
Documents / Settings (see build/CONTRACT.md). Nothing is computed in Python
except layout (row heights, the Checks last-row number).

Print layout (A4 landscape, 1 page wide, readable at >= 8 pt on paper):
  page 1  title, as-of, KPI tiles, checkpoint summary, legend
  page 2  per-flight readiness (22 flights)
  page 3  overdue actions (first 18)
  page 4  outstanding discrepancies (first 18)
The printable grid B..R is ~160 width units, so fit-to-width prints the
10 pt body at ~8.5 pt. List rows are fixed at 2 lines; long free text is
cut with an ellipsis (full text on the Checks sheet) so nothing is clipped.
"""
from openpyxl.formatting.rule import CellIsRule, DataBarRule, FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import column_index_from_string, get_column_letter
from openpyxl.worksheet.pagebreak import Break
from openpyxl.worksheet.properties import PageSetupProperties

SHEET = "Dashboard"
FONT = "Arial"
BODY = 10   # body text
NOTE = 9    # notes / captions

NAVY = "1F3864"
NAVY_MID = "2F5597"
NAVY_LIGHT = "D9E1F2"
GREY = "F2F2F2"
GREY_MID = "D9D9D9"
GREY_TEXT = "595959"
RED = "C00000"
RED_FILL = "F8CBAD"
GREEN_FILL = "C6EFCE"
GREEN_TEXT = "006100"
AMBER_FILL = "FFE699"
AMBER_TEXT = "7F6000"
BAR = "8EA9DB"

FIRST_FLIGHT_ROW, LAST_FLIGHT_ROW = 5, 26
CHECKPOINTS = [
    ("T-7D", "T-7D", "Preparation checks due 7 days before STD"),
    ("T-24H", "T-24H", "Preparation checks due 24 hours before STD"),
    ("T-12H PREP", "T-12H prep", "Preparation checks due 12 hours before STD"),
    ("UPLIFT", "Uplift", "Physical uplift / on-board confirmation, due at STD"),
]
TOP_OVERDUE = 18        # 2-line rows -> one printed page
TOP_DISCREPANCY = 18

# Printable grid B..R (A = unprinted margin, T = hidden helper). Total B..R = 163.
WIDTHS = {
    "A": 2, "B": 5, "C": 9, "D": 8, "E": 9, "F": 5, "G": 10, "H": 14,
    "I": 8, "J": 8, "K": 8, "L": 8, "M": 9, "N": 8, "O": 8, "P": 8,
    "Q": 8, "R": 30, "S": 2, "T": 8,
}
LAST_COL = "R"

thin = Side(style="thin", color=GREY_MID)
BORDER = Border(left=thin, right=thin, top=thin, bottom=thin)


def _font(size=BODY, bold=False, color="000000", italic=False):
    return Font(name=FONT, size=size, bold=bold, color=color, italic=italic)


def _fill(color):
    return PatternFill("solid", start_color=color, end_color=color)


def _span(a, b):
    return [get_column_letter(i) for i in
            range(column_index_from_string(a), column_index_from_string(b) + 1)]


def _width(a, b=None):
    return sum(WIDTHS[c] for c in _span(a, b or a))


def _lines(text, width, size=BODY, bold=False):
    """Conservative wrapped-line estimate for text in a cell `width` units wide."""
    per_unit = 1.1 * 10.0 / size * (0.9 if bold else 1.0)
    usable = max(width - 1.5, 1) * per_unit
    total = 0
    for para in str(text).split("\n"):
        # word wrap: a word never splits, so count greedily
        n, cur = 1, 0
        for w in para.split(" "):
            add = len(w) + (1 if cur else 0)
            if cur and cur + add > usable:
                n, cur = n + 1, len(w)
            else:
                cur += add
        total += n
    return total


def _chars(width, size=BODY):
    """Conservative characters per line for a cell `width` units wide."""
    return max(width - 1.5, 1) * 1.1 * 10.0 / size


def _height(lines, size=BODY):
    return round(lines * size * 1.2 + 4, 1)


class _Sheet:
    def __init__(self, ws):
        self.ws = ws

    def put(self, ref, value, *, size=BODY, bold=False, color="000000", italic=False,
            fill=None, h="left", v="center", wrap=False, fmt=None, border=False,
            merge_to=None):
        ws = self.ws
        c = ws[ref]
        c.value = value
        c.font = _font(size, bold, color, italic)
        c.alignment = Alignment(horizontal=h, vertical=v, wrap_text=wrap,
                                indent=(1 if h in ("left", "right") else 0))
        if fmt:
            c.number_format = fmt
        cells = [c]
        if merge_to:
            ws.merge_cells(f"{ref}:{merge_to}")
            cells = [x for row in ws[f"{ref}:{merge_to}"] for x in row]
        for x in cells:
            if fill:
                x.fill = _fill(fill)
            if border:
                x.border = BORDER
        return c

    def section(self, row, title, note=None):
        if note is None:
            self.put(f"B{row}", title, size=12, bold=True, color="FFFFFF", fill=NAVY_MID,
                     merge_to=f"{LAST_COL}{row}")
        else:  # title left, formula note right-aligned in the same band
            self.put(f"B{row}", title, size=12, bold=True, color="FFFFFF",
                     fill=NAVY_MID, merge_to=f"K{row}")
            self.put(f"L{row}", note, size=NOTE, italic=True, color="FFFFFF",
                     fill=NAVY_MID, h="right", merge_to=f"{LAST_COL}{row}")
        self.ws.row_dimensions[row].height = 22

    def header(self, row, spans):
        """spans: list of (first_col, last_col, text)."""
        n = 1
        for a, b, text in spans:
            self.put(f"{a}{row}", text, bold=True, color="FFFFFF", fill=NAVY,
                     h="center", wrap=True, border=True,
                     merge_to=(f"{b}{row}" if a != b else None))
            n = max(n, _lines(text, _width(a, b), bold=True))
        self.ws.row_dimensions[row].height = _height(n)


def build(wb, data):
    ws = wb.create_sheet(SHEET)
    S = _Sheet(ws)
    n_checks = len(data["checks"])
    CL = 5 + n_checks - 1  # last Checks data row

    def ck(col):
        return f"Checks!${col}$5:${col}${CL}"

    def fl(col):
        return f"Flights!${col}${FIRST_FLIGHT_ROW}:${col}${LAST_FLIGHT_ROW}"

    asof_note = '="As of "&TEXT(AsOfUTC+8/24,"dd-mmm-yy hh:mm")&" MYT"'

    ws.sheet_view.showGridLines = False
    ws.sheet_view.zoomScale = 90
    for col, w in WIDTHS.items():
        ws.column_dimensions[col].width = w
    ws.column_dimensions["T"].hidden = True

    # ============================================================ PAGE 1
    S.put("B1", "MAGCS Inflight Catering – Skytrax 2026 Readiness Dashboard",
          size=16, bold=True, color="FFFFFF", fill=NAVY, merge_to=f"{LAST_COL}1")
    ws.row_dimensions[1].height = 30
    S.put("B2", "As of (UTC)", bold=True, color=GREY_TEXT, fill=GREY, merge_to="C2")
    S.put("D2", "=AsOfUTC", bold=True, color=NAVY, fill=GREY, fmt="dd-mmm-yy hh:mm",
          merge_to="E2")
    S.put("F2", "As of (MYT)", bold=True, color=GREY_TEXT, fill=GREY, merge_to="G2")
    S.put("H2", "=AsOfUTC+8/24", bold=True, color=NAVY, fill=GREY, fmt="dd-mmm-yy hh:mm",
          merge_to="I2")
    note2 = ("Settings B4 (override) blank → live clock (recalculate to refresh); "
             "effective as-of shown in Settings B6.")
    S.put("J2", '=IF(Settings!$B$4<>"","Settings B4 (override) is set → as-of time is fixed",'
                '"Settings B4 (override) blank → live clock (recalculate to refresh)")'
                '&"; effective as-of shown in Settings B6."',
          size=NOTE, italic=True, color=GREY_TEXT, fill=GREY, wrap=True,
          merge_to=f"{LAST_COL}2")
    ws.row_dimensions[2].height = _height(_lines(note2, _width("J", LAST_COL), NOTE), NOTE)
    ws.row_dimensions[3].height = 6

    # ---------------- KPI tiles
    S.section(4, "Key indicators")
    tiles = [
        ("B", "C", "Flights", f'=COUNTA({fl("A")})', "0", "legs on the agenda", False),
        ("D", "E", "Flights READY", f'=COUNTIF({fl("AT")},"READY")', "0",
         "=\"of \"&COUNTA(" + fl("A") + ")&\" flights\"", False),
        ("F", "G", "Overall completion",
         f'=IF(SUM({ck("AI")})=0,"n/a",SUM({ck("AJ")})/SUM({ck("AI")}))',
         "0.0%", f'=SUM({ck("AJ")})&" of "&SUM({ck("AI")})&" checks"', False),
        ("H", "I", "Overdue actions", f'=SUM({ck("AK")})', "0", "past due, not complete", True),
        ("J", "K", "Open discrepancies", f'=SUM({ck("AL")})', "0", "fail, qty variance or CA open", True),
        ("L", "M", "Docs outstanding", f'=SUM({fl("AR")})', "0", "GLD + menu checklist", True),
        ("N", "O", "Clarifications open", f'=SUM({fl("AS")})', "0", "awaiting reference answer", True),
        ("P", "R", "Invalid entries", f'=SUM({ck("AM")})', "0", "entries failing validation rules", True),
    ]
    lab_lines = sub_lines = 1
    for a, b, label, formula, fmt, sub, alarm in tiles:
        S.put(f"{a}5", label, bold=True, color=NAVY, fill=NAVY_LIGHT, h="center", wrap=True,
              merge_to=f"{b}5")
        S.put(f"{a}6", formula, size=20, bold=True, color=NAVY, fill=GREY, h="center",
              fmt=fmt, merge_to=f"{b}6")
        S.put(f"{a}7", sub, size=NOTE, italic=True, color=GREY_TEXT, fill=GREY, h="center",
              v="top", wrap=True, merge_to=f"{b}7")
        lab_lines = max(lab_lines, _lines(label, _width(a, b), bold=True))
        sample = "9999 of 9999 checks" if sub.startswith("=") else sub
        sub_lines = max(sub_lines, _lines(sample, _width(a, b), NOTE))
        if alarm:
            ws.conditional_formatting.add(
                f"{a}6:{b}6", CellIsRule(operator="greaterThan", formula=["0"],
                                         font=Font(name=FONT, color=RED, bold=True)))
    ws.conditional_formatting.add(
        "F6:G6", FormulaRule(formula=['AND(ISNUMBER($F$6),$F$6>=1)'],
                             font=Font(name=FONT, color=GREEN_TEXT, bold=True)))
    ws.row_dimensions[5].height = _height(lab_lines)
    ws.row_dimensions[6].height = 32
    ws.row_dimensions[7].height = _height(sub_lines, NOTE)
    ws.row_dimensions[8].height = 8

    # ---------------- checkpoint summary
    r = 9
    S.section(r, "Checkpoint summary")
    r += 1
    S.header(r, [("B", "C", "Checkpoint"), ("D", "E", "In scope"), ("F", "G", "Complete"),
                 ("H", "H", "N/A (justified + rule)"), ("I", "I", "Open"),
                 ("J", "K", "Overdue"), ("L", "M", "Completion"),
                 ("N", LAST_COL, "What is checked")])
    first_cp = r + 1
    for key, label, desc in CHECKPOINTS:
        r += 1
        crit = f'{ck("G")},"{key}"'
        S.put(f"B{r}", label, bold=True, color=NAVY, border=True, merge_to=f"C{r}")
        S.put(f"D{r}", f"=SUMIFS({ck('AI')},{crit})", h="center", fmt="0", border=True,
              merge_to=f"E{r}")
        S.put(f"F{r}", f"=SUMIFS({ck('AJ')},{crit})", h="center", fmt="0", border=True,
              merge_to=f"G{r}")
        S.put(f"H{r}", f'=COUNTIFS({crit},{ck("AH")},"N/A \u2013 JUSTIFIED*")'
                       f'+COUNTIFS({crit},{ck("AH")},"N/A \u2013 RULE*")',
              h="center", fmt="0", border=True)
        S.put(f"I{r}", f"=D{r}-F{r}", h="center", fmt="0", border=True)
        S.put(f"J{r}", f"=SUMIFS({ck('AK')},{crit})", h="center", fmt="0", border=True,
              merge_to=f"K{r}")
        S.put(f"L{r}", f'=IF(D{r}=0,"n/a",F{r}/D{r})', h="center", fmt="0.0%", border=True,
              merge_to=f"M{r}")
        S.put(f"N{r}", desc, color=GREY_TEXT, border=True, merge_to=f"{LAST_COL}{r}")
        ws.row_dimensions[r].height = 18
    last_cp = r
    r += 1
    S.put(f"B{r}", "All", bold=True, color=NAVY, fill=NAVY_LIGHT, border=True,
          merge_to=f"C{r}")
    for a, b in (("D", "E"), ("F", "G"), ("H", "H"), ("I", "I"), ("J", "K")):
        S.put(f"{a}{r}", f"=SUM({a}{first_cp}:{a}{last_cp})", bold=True, h="center",
              fmt="0", fill=NAVY_LIGHT, border=True, merge_to=(f"{b}{r}" if a != b else None))
    S.put(f"L{r}", f'=IF(D{r}=0,"n/a",F{r}/D{r})', bold=True, h="center", fmt="0.0%",
          fill=NAVY_LIGHT, border=True, merge_to=f"M{r}")
    S.put(f"N{r}", "Open = in scope \u2212 complete (includes overdue).", italic=True,
          color=GREY_TEXT, fill=NAVY_LIGHT, border=True, merge_to=f"{LAST_COL}{r}")
    ws.row_dimensions[r].height = 18
    ws.conditional_formatting.add(
        f"L{first_cp}:L{r}", DataBarRule(start_type="num", start_value=0, end_type="num",
                                         end_value=1, color=BAR, showValue=True))
    ws.conditional_formatting.add(
        f"J{first_cp}:J{r}", CellIsRule(operator="greaterThan", formula=["0"],
                                        font=Font(name=FONT, color=RED, bold=True)))
    r += 1
    ws.row_dimensions[r].height = 8

    # ---------------- legend
    r += 1
    S.section(r, "How the figures are calculated")
    legend = [
        ("Blank checks", "A check line with no Status / completion entry never counts as complete. "
                         "It stays OPEN until its due time, then becomes OVERDUE."),
        ("N/A", "N/A only removes a check from the denominator when it is either rule-based "
                "(applicability “N/A – rule”, e.g. fleet-specific items) or JUSTIFIED: "
                "Status N/A with a written N/A justification AND a named verifier. An unjustified "
                "or unverified N/A is flagged INVALID and still counts as open."),
        ("Complete", "Status Pass or Pass after CA, with a completion time, evidence ref and verifier; "
                     "plus the batch ID on T-24H lines and expected + actual quantity on quantity "
                     "lines. A quantity variance is only accepted as Pass after CA with the CA status "
                     "Closed. GLD / menu-checklist lines also need the document ON FILE (Documents "
                     "sheet). The completion time must not be in the future or before the valid "
                     "window, and uplift lines can only be completed within the uplift window before "
                     "STD (Settings B7). Anything else is OPEN, OVERDUE or INVALID."),
        ("Discrepancy", "An in-scope check is an open discrepancy when its Status is Fail, OR its CA "
                        "status is Open, OR its quantity variance is not zero (unless Status is Pass "
                        "after CA). It clears once the corrective action is recorded and closed."),
        ("READY", "A flight is READY only when every in-scope check at all four checkpoints is "
                  "complete – including the physical uplift confirmation – with zero open "
                  "discrepancies, the GLD and menu checklist on file (Documents sheet) and no invalid "
                  "entries. PREP DONE – AWAITING UPLIFT means all preparation checks are complete "
                  "but the physical uplift is not yet confirmed."),
        ("Lists", f"The overdue and discrepancy pages list the first {TOP_OVERDUE} / "
                  f"{TOP_DISCREPANCY} lines in Checks order. Long text there is shortened with "
                  "\u201c\u2026\u201d \u2013 the full text is on the Checks sheet."),
        ("Time basis", "Due times are computed in UTC from STD and shown in local time at the check "
                       "station. As-of time is Settings B6: the override in B4 if set, otherwise the "
                       "PC clock adjusted by the offset in B5."),
    ]
    for label, text in legend:
        r += 1
        S.put(f"B{r}", label, bold=True, color=NAVY, fill=NAVY_LIGHT, v="top", border=True,
              wrap=True, merge_to=f"D{r}")
        S.put(f"E{r}", text, v="top", wrap=True, border=True, merge_to=f"{LAST_COL}{r}")
        ws.row_dimensions[r].height = _height(max(_lines(text, _width("E", LAST_COL)),
                                                  _lines(label, _width("B", "D"), bold=True)))

    # ============================================================ PAGE 2: per-flight
    r += 1
    ws.row_breaks.append(Break(id=r))
    r += 1
    S.section(r, "Per-flight readiness", note=asof_note)
    r += 1
    cols = [
        ("B", "Itin", "B", "0"), ("C", "Flight", "F", None), ("D", "Date", "D", "dd-mmm"),
        ("E", "Sector", None, None), ("F", "Cls", "K", None), ("G", "Fleet", "L", None),
        ("H", "STD local", "I", "dd-mmm hh:mm"),
        ("I", "T-7D", "AJ", "0%"), ("J", "T-24H", "AK", "0%"),
        ("K", "T-12H prep", "AL", "0%"), ("L", "Uplift", "AM", "0%"),
        ("M", "Overall", "AN", "0%"), ("N", "Over-due", "AP", "0"),
        ("O", "Dis-crep.", "AQ", "0"), ("P", "Docs out", "AR", "0"),
        ("Q", "Clarif. open", "AS", "0"), ("R", "Readiness", "AT", None),
    ]
    S.header(r, [(c, c, t) for c, t, _, _ in cols])
    f0 = r + 1
    for i in range(LAST_FLIGHT_ROW - FIRST_FLIGHT_ROW + 1):
        r += 1
        src = FIRST_FLIGHT_ROW + i
        band = GREY if i % 2 else None
        for c, _t, fcol, fmt in cols:
            if fcol is None:  # sector = Dep-Arr
                val = f'=IF(Flights!$G${src}="","",Flights!$G${src}&"-"&Flights!$H${src})'
            else:
                ref = f"Flights!${fcol}${src}"
                val = f'=IF({ref}="","",{ref})'
            S.put(f"{c}{r}", val, h=("left" if c in ("C", "R") else "center"),
                  size=(NOTE if c == "R" else BODY),  # longest state = 27 chars
                  bold=(c in ("C", "R")), fmt=fmt, fill=band, border=True)
        ws.row_dimensions[r].height = 19
    f1 = r
    ws.conditional_formatting.add(
        f"I{f0}:M{f1}", DataBarRule(start_type="num", start_value=0, end_type="num",
                                    end_value=1, color=BAR, showValue=True))
    ws.conditional_formatting.add(
        f"N{f0}:Q{f1}", CellIsRule(operator="greaterThan", formula=["0"],
                                   font=Font(name=FONT, color=RED, bold=True)))
    rd = f"R{f0}:R{f1}"
    ws.conditional_formatting.add(rd, FormulaRule(
        formula=[f'$R{f0}="READY"'], fill=_fill(GREEN_FILL),
        font=Font(name=FONT, color=GREEN_TEXT, bold=True), stopIfTrue=True))
    ws.conditional_formatting.add(rd, FormulaRule(
        formula=[f'LEFT($R{f0},9)="NOT READY"'], fill=_fill(RED_FILL),
        font=Font(name=FONT, color=RED, bold=True), stopIfTrue=True))
    ws.conditional_formatting.add(rd, FormulaRule(
        formula=[f'LEFT($R{f0},9)="PREP DONE"'], fill=_fill(AMBER_FILL),
        font=Font(name=FONT, color=AMBER_TEXT, bold=True), stopIfTrue=True))
    ws.conditional_formatting.add(rd, FormulaRule(
        formula=[f'OR($R{f0}="IN PROGRESS",$R{f0}="NOT STARTED")'], fill=_fill(GREY_MID),
        font=Font(name=FONT, color=GREY_TEXT, bold=True), stopIfTrue=True))
    r += 1
    fnote = ("T-7D / T-24H / T-12H prep / Uplift / Overall = % of in-scope checks complete "
             "(“n/a” = no in-scope checks). Over-due, Dis-crep., Docs out, Clarif. open "
             "are counts; above zero shown in red. Source: Flights sheet rows 5–26.")
    S.put(f"B{r}", fnote, size=NOTE, italic=True, color=GREY_TEXT, wrap=True, v="top",
          merge_to=f"{LAST_COL}{r}")
    ws.row_dimensions[r].height = _height(_lines(fnote, _width("B", LAST_COL), NOTE), NOTE)

    # ============================================================ PAGES 3-4: lists
    def top_list(r, title, top_n, total_expr, seq_col, spans, noun, lines):
        """spans: (first_col, last_col, header, checks_col, kind)."""
        S.section(r, title,
                  note=f'="Showing "&MIN({top_n},{total_expr})&" of "&{total_expr}'
                       f'&"  •  as of "&TEXT(AsOfUTC+8/24,"dd-mmm-yy hh:mm")&" MYT"')
        r += 1
        S.header(r, [("B", "B", "#")] + [(a, b, h) for a, b, h, _, _ in spans])
        ws[f"T{r}"].value = "Checks row"
        ws[f"T{r}"].font = _font(8, color=GREY_TEXT)
        for k in range(1, top_n + 1):
            r += 1
            band = GREY if k % 2 == 0 else None
            ws[f"T{r}"].value = f"=IFERROR(MATCH({k},{ck(seq_col)},0)+4,\"\")"  # sheet row (data start row 5)
            ws[f"T{r}"].font = _font(8, color=GREY_TEXT)
            S.put(f"B{r}", f'=IF($T{r}="","",{k})', size=NOTE, color=GREY_TEXT, h="center",
                  v="top", fill=band, border=True)
            for a, b, _h, ccol, kind in spans:
                cap = max(8, int(_chars(_width(a, b)) * lines * 0.85))
                if kind == "duestn":  # due local + check station, one cell
                    val = (f'=IF($T{r}="","",IF(ISNUMBER(INDEX({ck("R")},$T{r}-4)),'
                           f'TEXT(INDEX({ck("R")},$T{r}-4),"dd-mmm hh:mm"),"")'
                           f'&" "&INDEX({ck("P")},$T{r}-4))')
                elif kind == "num":
                    rng = ck(ccol)
                    val = (f'=IF($T{r}="","",IF(INDEX({rng},$T{r}-4)="","",'
                           f'INDEX({rng},$T{r}-4)))')
                else:  # text: cut to what fits in `lines` lines, with an ellipsis
                    x = f"INDEX({ck(ccol)},$T{r}-4)"
                    val = (f'=IF($T{r}="","",IF(LEN({x})>{cap},LEFT({x},{cap - 1})&"\u2026",'
                           f'{x}&""))')
                S.put(f"{a}{r}", val,
                      h=("left" if (kind == "text" and _width(a, b) > 12) else "center"),
                      v="top", wrap=True, fmt=("#,##0;-#,##0;0" if kind == "num" else None),
                      fill=band, border=True, merge_to=(f"{b}{r}" if a != b else None))
            ws.row_dimensions[r].height = _height(lines)
        r += 1
        S.put(f"B{r}", f'=IF({total_expr}=0,"No {noun} at the as-of time.",'
                       f'IF({total_expr}>{top_n},"+"&({total_expr}-{top_n})&" more {noun} '
                       f'not shown – see Checks sheet (filter on column {seq_col}).",""))',
              size=NOTE, bold=True, italic=True, color=GREEN_TEXT, merge_to=f"{LAST_COL}{r}")
        ws.conditional_formatting.add(f"B{r}", FormulaRule(
            formula=[f'LEFT($B${r},1)="+"'], font=Font(name=FONT, color=RED, bold=True)))
        ws.row_dimensions[r].height = 16
        return r

    # Overdue actions
    od_spans = [("C", "D", "CheckID", "A", "text"), ("E", "E", "Flight", "C", "text"),
                ("F", "G", "Checkpoint", "G", "text"), ("H", "M", "Check item", "J", "text"),
                ("N", "O", "Due local @ stn", None, "duestn"), ("P", "Q", "PIC", "T", "text"),
                ("R", "R", "Record state", "AH", "text")]
    r += 1
    ws.row_breaks.append(Break(id=r))
    r += 1
    r = top_list(r, f"Overdue actions (first {TOP_OVERDUE})", TOP_OVERDUE,
                 f"SUM({ck('AK')})", "AO", od_spans, "overdue actions", 2)

    # Outstanding discrepancies
    ds_spans = [("C", "D", "CheckID", "A", "text"), ("E", "E", "Flight", "C", "text"),
                ("F", "G", "Checkpoint", "G", "text"), ("H", "K", "Check item", "J", "text"),
                ("L", "L", "Status", "U", "text"), ("M", "M", "Exp. qty", "X", "num"),
                ("N", "N", "Act. qty", "Y", "num"), ("O", "O", "Vari-ance", "AF", "num"),
                ("P", "P", "CA status", "AB", "text"), ("Q", "R", "Corrective action", "AA", "text")]
    r += 1
    ws.row_breaks.append(Break(id=r))
    r += 1
    r = top_list(r, f"Outstanding discrepancies (first {TOP_DISCREPANCY})", TOP_DISCREPANCY,
                 f"SUM({ck('AL')})", "AP", ds_spans, "outstanding discrepancies", 2)
    last_row = r

    # ============================================================ view / print
    ws.freeze_panes = "A4"
    ws.print_area = f"B1:{LAST_COL}{last_row}"
    ws.page_setup.orientation = "landscape"
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr = PageSetupProperties(fitToPage=True)
    ws.print_options.horizontalCentered = True
    ws.page_margins.left = ws.page_margins.right = 0.25
    ws.page_margins.top = 0.35
    ws.page_margins.bottom = 0.45
    ws.page_margins.header = 0.15
    ws.page_margins.footer = 0.15
    ws.oddFooter.left.text = "MAGCS – Skytrax 2026 catering readiness – Dashboard"
    ws.oddFooter.left.size = 8
    ws.oddFooter.right.text = "Page &P of &N"
    ws.oddFooter.right.size = 8
    ws.sheet_properties.tabColor = NAVY
    return ws

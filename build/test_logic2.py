"""Regression tests for logic-critic round 2 findings."""
import shutil, subprocess, sys
from datetime import datetime, timedelta
import openpyxl
sys.path.insert(0, "/home/user/skytrax-checklist/build")
import test_ready as T

F11_STD = datetime(2026, 10, 12, 15, 50)  # PEN local


def fill_f11(wb):
    T.STD = F11_STD
    T.fill(wb["Checks"], wb["Documents"], fid="F11")
    ck = wb["Checks"]
    for r in range(5, ck.max_row + 1):
        if ck[f"B{r}"].value != "F11":
            continue
        cid = ck[f"A{r}"].value
        if cid == "F11-T7-10":  # carrying flight clarification: must be Pass (N/A not allowed)
            ck[f"U{r}"] = "Pass"; ck[f"AC{r}"] = "Confirmed – applies / carried as listed"
        if cid == "F11-UPL-03":  # KUL loading line: before MH1140 dep 11:45 KUL
            ck[f"AD{r}"] = datetime(2026, 10, 12, 10, 0)
        if ck[f"G{r}"].value in ("T-24H", "T-12H PREP"):  # capped by loading (AY = 12-Oct 05:45 KUL)
            ck[f"AD{r}"] = datetime(2026, 10, 11, 23, 0) if ck[f"G{r}"].value == "T-12H PREP" else datetime(2026, 10, 11, 10, 0)
        if ck[f"H{r}"].value == "Physical uplift" and cid != "F11-UPL-03":
            ck[f"AD{r}"] = datetime(2026, 10, 12, 14, 0)  # PEN, after MH1140 arrival
    wb["Flights"]["AX15"] = datetime(2026, 10, 12, 3, 45)


def run(tag, mutate):
    p = T.TMP.format("l2_" + tag)
    shutil.copy(T.SRC, p)
    wb = openpyxl.load_workbook(p)
    wb["Settings"]["B4"] = datetime(2026, 10, 12, 7, 49)
    fill_f11(wb)
    rows = {wb["Checks"].cell(r, 1).value: r for r in range(5, wb["Checks"].max_row + 1)}
    if mutate:
        mutate(wb, rows)
    wb.save(p)
    out = subprocess.run([sys.executable, T.RECALC, p, "900"], capture_output=True, text=True).stdout
    wb = openpyxl.load_workbook(p, data_only=True)
    ck = wb["Checks"]
    bad = [(ck[f"A{r}"].value, ck[f"AH{r}"].value) for r in range(5, ck.max_row + 1)
           if ck[f"B{r}"].value == "F11" and not str(ck[f"AH{r}"].value).startswith(("COMPLETE", "N/A"))]
    ok = '"total_errors": 0' in out
    print(f"{tag:26s} {wb['Flights']['AT15'].value!r:32s} ok={ok} bad={bad[:3]}")


def cell(cid, col, v):
    return lambda wb, rows: wb["Checks"].__setitem__(f"{col}{rows[cid]}", v)


if __name__ == "__main__" and len(sys.argv) == 1:
  run("baseline_F11", None)
  run("onboard_before_carry", cell("F11-UPL-01", "AD", datetime(2026, 10, 12, 10, 0)))
  run("ca_open_space", lambda wb, rows: [wb["Checks"].__setitem__(f"AA{rows['F11-T7-01']}", "Re-plate"),
                                         wb["Checks"].__setitem__(f"AB{rows['F11-T7-01']}", "Open ")])
  run("carry_clar_NA", lambda wb, rows: [wb["Checks"].__setitem__(f"U{rows['F11-T7-10']}", "N/A"),
                                         wb["Checks"].__setitem__(f"AC{rows['F11-T7-10']}", "n/a")])
  run("ax_after_std", lambda wb, rows: wb["Flights"].__setitem__("AX15", datetime(2026, 10, 12, 9, 0)))
  run("result_blank", cell("F11-T24-01", "V", None))
  run("t7_stale", cell("F11-T7-01", "AD", datetime(2026, 7, 2, 10, 0)))
  run("pic_is_verifier", cell("F11-T7-01", "AE", "A. Rahman"))
  run("ca_text_no_status", cell("F11-T7-02", "AA", "Reprinted cards"))


def more():
    run("ok_evidence_cells_ignored", cell("F11-T7-01", "Z", "TBC"))  # S / Z no longer used
    run("placeholder_verifier", cell("F11-T7-01", "AE", "?"))
    run("prep_after_carrier", cell("F11-T12-01", "AD", datetime(2026, 10, 12, 12, 0)))
    run("onboard_before_arrival", cell("F11-UPL-01", "AD", datetime(2026, 10, 12, 12, 30)))
    run("pass_after_ca_short", lambda wb, rows: [wb["Checks"].__setitem__(f"{k}{rows['F11-UPL-01']}", v) for k, v in
                                                 (("X", 20), ("Y", 15), ("U", "Pass after CA"), ("AA", "Topped up 5 meals from spare"), ("AB", "Closed"))])
    run("na_short_just", lambda wb, rows: [wb["Checks"].__setitem__(f"U{rows['F11-T7-02']}", "N/A"),
                                           wb["Checks"].__setitem__(f"AC{rows['F11-T7-02']}", "x")])
    run("uplift_window_24", lambda wb, rows: wb["Settings"].__setitem__("B7", 24))


if __name__ == "__main__" and len(sys.argv) > 1 and sys.argv[1] == "more":
    more()


def round4():
    d = lambda k, v: (lambda wb, rows: wb["Documents"].__setitem__(k, v))
    run("doc_placeholders", lambda wb, rows: [wb["Documents"].__setitem__(k, v) for k, v in (("F15", "TBC"), ("H15", "pending"))])
    run("verifier_same_person", cell("F11-T7-01", "AE", "A Rahman"))
    run("expected_lowered", lambda wb, rows: [wb["Checks"].__setitem__(f"{k}{rows['F11-UPL-01']}", v) for k, v in (("X", 8), ("Y", 8))])
    run("pass_but_rejected", cell("F11-T24-01", "V", "Rejected - off taste, batch discarded"))
    run("error_value_pasted", cell("F11-T7-01", "V", "#N/A"))  # recalc reports the pasted error itself (ok=False expected)


if __name__ == "__main__" and len(sys.argv) > 1 and sys.argv[1] == "r4":
    round4()


def round5():
    run("result_discrepancy_pass", cell("F11-T12-03", "V", "Discrepancy found - 3 dirty inserts replaced"))
    run("ok_result_shortbread", cell("F11-T24-01", "V", "Shortbread texture and taste good, temp 4C"))


if __name__ == "__main__" and len(sys.argv) > 1 and sys.argv[1] == "r5":
    round5()


def round6():
    ok_results = ["No defects found", "No discrepancies", "Failsafe seal intact", "Nil discrepancy", "no stains or defects",
                  "Shortbread and dessert to spec", "Defect-free, temps 4C"]
    for i, (cid, txt) in enumerate(zip(["F11-T7-01", "F11-T7-02", "F11-T7-04", "F11-T7-05", "F11-T7-06", "F11-T24-01", "F11-T12-01"], ok_results)):
        run(f"ok_result_{i}", cell(cid, "V", txt))
    run("bad_result_awaiting", cell("F11-T24-01", "V", "Awaiting panel score"))
    run("bad_result_short", cell("F11-T12-03", "V", "Short 3 trays, replaced"))
    run("bad_batch_tbc", cell("F11-T24-01", "W", "Batch TBC"))
    run("bad_doc_to_follow", lambda wb, rows: wb["Documents"].__setitem__("F15", "to follow from PASB"))
    run("load_change_ca", lambda wb, rows: [wb["Checks"].__setitem__(f"{k}{rows['F11-UPL-01']}", v) for k, v in
                                            (("X", 14), ("Y", 14), ("U", "Pass after CA"),
                                             ("AA", "Final load +2 pax: 2 extra refreshments uplifted"), ("AB", "Closed"))])


if __name__ == "__main__" and len(sys.argv) > 1 and sys.argv[1] == "r6":
    round6()


def round7():
    ok = {"F11-T7-01": "0 discrepancies", "F11-T7-02": "Discrepancies: 0", "F11-T7-04": "Clean - 0 defects",
          "F11-T7-05": "fail-safe seals intact", "F11-T7-06": "Nothing missing, count matches",
          "F11-T24-01": "Short rib and rice to spec", "F11-T12-01": "Menu cards correct, not damaged"}
    for i, (cid, txt) in enumerate(ok.items()):
        run(f"r7_ok_result_{i}", cell(cid, "V", txt))
    run("r7_ok_na_not_available", lambda wb, rows: [wb["Checks"].__setitem__(f"U{rows['F11-T7-02']}", "N/A"),
        wb["Checks"].__setitem__(f"AC{rows['F11-T7-02']}", "Refreshment service – no printed menu card")])
    run("r7_bad_missing_no_spares", cell("F11-T12-03", "V", "2 BC meals missing, no spares available"))
    run("r7_bad_verifier_self", cell("F11-T7-01", "AE", "Self"))
    run("r7_bad_prep_same_minute_as_loading", cell("F11-T12-01", "AD", datetime(2026, 10, 12, 14, 0)))


if __name__ == "__main__" and len(sys.argv) > 1 and sys.argv[1] == "r7":
    round7()


def round17():
    run("r17_bad_not_delivered", cell("F11-T12-01", "V", "Menu cards not delivered by caterer"))
    run("r17_bad_meals_short", cell("F11-UPL-01", "V", "2 meals short, loaded 22 of 24"))
    run("r17_ok_short_rib", cell("F11-T24-01", "V", "Short rib and rice to spec, nothing short"))
    run("r17_bad_na_filler", lambda wb, rows: [wb["Checks"].__setitem__(f"U{rows['F11-T7-02']}", "N/A"),
        wb["Checks"].__setitem__(f"AC{rows['F11-T7-02']}", "Not applicable for this flight")])
    run("r17_bad_na_as_above", lambda wb, rows: [wb["Checks"].__setitem__(f"U{rows['F11-T7-02']}", "N/A"),
        wb["Checks"].__setitem__(f"AC{rows['F11-T7-02']}", "As above - see previous line")])
    run("r17_bad_fail_on_rule_row", lambda wb, rows: [wb["Checks"].__setitem__(f"U{rows['F11-T7-03']}", "Fail"),
        wb["Checks"].__setitem__(f"AA{rows['F11-T7-03']}", "n/a"), wb["Checks"].__setitem__(f"AB{rows['F11-T7-03']}", "Open")])


if __name__ == "__main__" and len(sys.argv) > 1 and sys.argv[1] == "r17":
    round17()


def round18():
    C = lambda cid, kv: (lambda wb, rows: [wb["Checks"].__setitem__(f"{k}{rows[cid]}", v) for k, v in kv])
    # (clarification-line outcome tests retired: MAGCS resolved every clarification on 01-Oct-2026)
    run("r18_bad_na_wrong_outcome", C("F11-T12-08", [("AC", "Refreshment service – no printed menu card")]))
    run("r18_bad_pass_with_na_outcome", C("F11-T7-01", [("AC", "Confirmed – not carried on this sector")]))
    # wording safety net
    for i, txt in enumerate(["Partially loaded", "Loading not complete", "Only 40 of 44 slippers on board",
                             "Juice out of date", "None missing except 2 meals", "Hot meals uplifted but no cutlery"]):
        run(f"r18_bad_wording_{i}", cell("F11-UPL-01", "V", txt))
    for i, txt in enumerate(["Com tam (broken rice) to spec, panel of 3", "Dirty linen bags positioned at G2 per GLD",
                             "Late-night supper service items all to spec"]):
        run(f"r18_ok_wording_{i}", cell("F11-T24-01", "V", txt))
    for i, txt in enumerate(["Pass", "All good", "As per menu", "Conforms"]):
        run(f"r18_bad_brief_{i}", cell("F11-T7-01", "V", txt))
    run("r18_bad_verifier_fa", cell("F11-T7-01", "AE", "Flight attendant"))
    run("r18_bad_verifier_csm1", cell("F11-T7-01", "AE", "CSM1"))
    run("r18_bad_same_initial", cell("F11-T7-01", "AE", "A. Rahman") if False else
        (lambda wb, rows: [wb["Checks"].__setitem__(f"T{rows['F11-T7-01']}", "Aisyah Rahman"),
                           wb["Checks"].__setitem__(f"AE{rows['F11-T7-01']}", "A. Rahman")]))
    run("r18_ok_patronymic", lambda wb, rows: [wb["Checks"].__setitem__(f"T{rows['F11-T7-01']}", "Ahmad bin Ali"),
                                               wb["Checks"].__setitem__(f"AE{rows['F11-T7-01']}", "Ali bin Ahmad")])
    run("r18_ok_doc_binder", lambda wb, rows: wb["Documents"].__setitem__("F15", "Filed in GLD binder, PEN catering office"))


if __name__ == "__main__" and len(sys.argv) > 1 and sys.argv[1] == "r18":
    round18()


def round19():
    for i, (pic, ver) in enumerate([("Raj Kumar", "Ravi Kumar"), ("John Smith", "Jane Smith")]):
        run(f"r19_ok_people_{i}", lambda wb, rows, pic=pic, ver=ver: [wb["Checks"].__setitem__(f"T{rows['F11-T7-01']}", pic),
                                                                   wb["Checks"].__setitem__(f"AE{rows['F11-T7-01']}", ver)])
    for i, txt in enumerate(["Not all meals loaded", "Meal temperature too high at 12C", "Toiletry kits left at caterer"]):
        run(f"r19_bad_wording_{i}", cell("F11-UPL-01", "V", txt))
    for i, txt in enumerate(["Leak test on water bottles passed, seals intact", "Missing items: none, all 12 on board",
                             "Damaged-cart log reviewed, zero damaged carts"]):
        run(f"r19_ok_wording_{i}", cell("F11-UPL-01", "V", txt))


if __name__ == "__main__" and len(sys.argv) > 1 and sys.argv[1] == "r19":
    round19()


def round24():
    for i, txt in enumerate(["None on board", "Nothing loaded", "Toiletry kits absent", "Cart not seen on board"]):
        run(f"r24_bad_onboard_{i}", cell("F11-UPL-01", "V", txt))
    for i, txt in enumerate(["No item shortfall found", "Not a single defect found",
                             "Special meals loaded after late pax update, all on board", "Wet towels loaded, stains absent"]):
        run(f"r24_ok_wording_{i}", cell("F11-UPL-01", "V", txt))


if __name__ == "__main__" and len(sys.argv) > 1 and sys.argv[1] == "r24":
    round24()

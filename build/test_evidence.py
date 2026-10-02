"""PIC / verifier name rules on F11 (evidence fields were removed at MAGCS request, 02-Oct-2026)."""
import sys, shutil, subprocess
from datetime import datetime
import openpyxl
sys.argv = [sys.argv[0]]
sys.path.insert(0, "/home/user/skytrax-checklist/build")
import test_logic2 as L, test_ready as T

IDS = ["T7-01", "T7-02", "T7-04", "T7-05", "T7-06", "T12-01", "T12-02", "T12-03", "T12-04", "T12-05", "T12-06", "T12-07"]


def run_verifiers():
    cases = ["Siti Aminah", "QA", "CSM", "Duty CSM", "A. Rahman", "Rahman", "PASB supervisor"]
    p = T.TMP.format("ev_verifier"); shutil.copy(T.SRC, p)
    wb = openpyxl.load_workbook(p); wb["Settings"]["B4"] = datetime(2026, 10, 12, 7, 49); L.fill_f11(wb)
    ck = wb["Checks"]; rows = {ck.cell(r, 1).value: r for r in range(5, ck.max_row + 1)}
    for sid, v in zip(IDS, cases):
        ck[f"AE{rows['F11-' + sid]}"] = v
    wb.save(p); subprocess.run([sys.executable, T.RECALC, p, "900"], capture_output=True)
    ck = openpyxl.load_workbook(p, data_only=True)["Checks"]
    for sid, v in zip(IDS, cases):
        print(f"verif {v:16s} (PIC {ck['T' + str(rows['F11-' + sid])].value}) -> {ck['AH' + str(rows['F11-' + sid])].value[:70]}")



def run_people():
    """(PIC, verifier, expected ok) on 12 rows."""
    cases = [("A. Rahman", "Siti Aminah", True), ("Toby Lim", "Ruby Tan", True), ("Esra Yilmaz", "Isra Omar", True),
             ("Abby Wong", "Ali", True), ("A. Rahman", "Shift Supervisor", False), ("A. Rahman", "Station Manager KUL", False),
             ("A. Rahman", "Purser", False), ("A. Rahman", "Quality Assurance", False), ("A. Rahman", "Cabin Services Manager", False),
             ("Aisyah Rahman", "Rahman, Aisyah", False), ("Station Manager KUL", "Siti Aminah", False), ("A. Rahman", "SATS QA", False)]
    p = T.TMP.format("ev_people"); shutil.copy(T.SRC, p)
    wb = openpyxl.load_workbook(p); wb["Settings"]["B4"] = datetime(2026, 10, 12, 7, 49); L.fill_f11(wb)
    ck = wb["Checks"]; rows = {ck.cell(r, 1).value: r for r in range(5, ck.max_row + 1)}
    for sid, (pic, ver, _) in zip(IDS, cases):
        ck[f"T{rows['F11-' + sid]}"] = pic; ck[f"AE{rows['F11-' + sid]}"] = ver
    wb.save(p); subprocess.run([sys.executable, T.RECALC, p, "900"], capture_output=True)
    ck = openpyxl.load_workbook(p, data_only=True)["Checks"]
    for sid, (pic, ver, exp) in zip(IDS, cases):
        st = ck['AH' + str(rows['F11-' + sid])].value
        flag = "ok " if st.startswith("COMPLETE") == exp else "BAD"
        print(f"people {flag} {pic:20s} / {ver:22s} -> {st[:70]}")


if __name__ == "__main__":
    run_verifiers()
    run_people()
